"""Behaviour tests for the one-off Zava migration."""

import importlib.util
from io import BytesIO
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

SPEC = importlib.util.spec_from_file_location("rename_zava", Path(__file__).with_name("rename-zava.py"))
rename = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(rename)


class RenameTests(unittest.TestCase):
    def test_display_and_technical_forms(self):
        self.assertEqual(
            rename.replace_names("Proper North Homes Proper Warm ProperWarm properwarm_mvp PW_AUTH pw_session X-PW-CSRF"),
            "Zava North Homes Zava Forward ZavaForward zavaforward_mvp ZF_AUTH zf_session X-ZF-CSRF",
        )

    def test_historical_external_urls_remain_exact(self):
        url = "https://github.com/jmservera/ama/blob/abc/PW/properwarm_mvp/server.py"
        self.assertEqual(rename.replace_names(f"Proper Warm {url}"), f"Zava Forward {url}")

    def test_official_warm_homes_funding_is_not_rebranded(self):
        self.assertEqual(rename.replace_names("Warm Homes: Social Housing Fund"), "Warm Homes: Social Housing Fund")

    def test_history_exceptions_and_active_ppt_sources(self):
        for path in ("docs/reference/log.md", "docs/reference/memory.md",
                     "docs/reference/regulatory/source.md", ".copilot-tracking/changes/old.md"):
            self.assertTrue(rename.preserved(path))
        for path in ("docs/prds/proper-warm.md", ".copilot-tracking/ppt/a/content/slide-001/content.yaml",
                     ".copilot-tracking/ppt/a/slide-deck/proper-warm.pptx"):
            self.assertFalse(rename.preserved(path))

    def test_paths_and_preserved_relative_references(self):
        history = ".copilot-tracking/changes/proper-warm.md"
        mapping = {history: history, "mvp/properwarm_mvp/a.py": "mvp/zavaforward_mvp/a.py"}
        result = rename.update_text(f"../{history} mvp/properwarm_mvp/a.py", "docs/example.md", mapping)
        self.assertEqual(result, f"../{history} mvp/zavaforward_mvp/a.py")

    def test_relative_link_into_historical_tracking_record(self):
        history = ".copilot-tracking/changes/proper-warm.md"
        result = rename.update_text(
            "[evidence](../../../.copilot-tracking/changes/proper-warm.md)",
            "docs/reference/research/page.md", {history: history},
        )
        self.assertIn("changes/proper-warm.md", result)

    def test_yaml_logo_source_and_dimensions(self):
        from PIL import Image
        import yaml
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "content/slide-001/content.yaml"
            path.parent.mkdir(parents=True)
            style = root / "content/global/style.yaml"
            style.parent.mkdir()
            style.write_text("dimensions:\n  width_inches: 13.333\n")
            logo = BytesIO()
            Image.new("RGB", (600, 200)).save(logo, "PNG")
            result = rename.add_yaml_logo("elements: []\n", path, logo=logo.getvalue(), apply=True)
            data = yaml.safe_load(result)
            self.assertEqual(data["elements"][-1]["height"], .2833)
            self.assertEqual((path.parent / "images/zava-logo.png").read_bytes(), logo.getvalue())
            self.assertEqual(rename.add_yaml_logo(result, path, logo=logo.getvalue(), apply=False), result)

    def test_style_rename_does_not_break_later_slide_sources(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", root], check=True)
            content = root / ".copilot-tracking/ppt/date/proper-warm/content"
            (content / "global").mkdir(parents=True)
            (content / "slide-001").mkdir()
            (content / "global/style.yaml").write_text("dimensions:\n  width_inches: 13.333\n")
            (content / "slide-001/content.yaml").write_text("elements: []\n")
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            logo = BytesIO()
            Image.new("RGB", (644, 225)).save(logo, "PNG")
            report = rename.migrate(root, apply=True, logo=logo.getvalue())
            self.assertEqual(report["failures"], [])
            new = root / ".copilot-tracking/ppt/date/zava-forward/content"
            self.assertTrue((new / "global/style.yaml").exists())
            self.assertTrue((new / "slide-001/images/zava-logo.png").exists())
            self.assertIn("Zava branding logo", (new / "slide-001/content.yaml").read_text())

    def test_policy_metadata_rename_recomputes_integrity_hash_without_changing_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", root], check=True)
            path = root / "mvp/policies/ranking-v1.json"
            path.parent.mkdir(parents=True)
            rules = [{"code": "waiting", "label": "Longest waiting first"}]
            path.write_text(json.dumps({"policy_id": "proper-warm-home-ranking",
                                       "rules": rules, "content_hash": "old"}))
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            report = rename.migrate(root, apply=True, logo=None)
            self.assertFalse(report["failures"])
            policy = json.loads(path.read_text())
            self.assertEqual(policy["rules"], rules)
            body = {k: v for k, v in policy.items() if k != "content_hash"}
            self.assertEqual(policy["content_hash"], hashlib.sha256(json.dumps(
                body, sort_keys=True, separators=(",", ":")).encode()).hexdigest())

    def test_historical_links_keep_exact_original_snapshot(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", root], check=True)
            document = root / "docs/prds/proper-warm.md"
            document.parent.mkdir(parents=True)
            document.write_text("Proper Warm original evidence")
            log = root / "docs/reference/log.md"
            log.parent.mkdir()
            log.write_text("[old evidence](../prds/proper-warm.md)")
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            report = rename.migrate(root, apply=True, logo=None)
            self.assertEqual(report["historical_snapshots"], ["docs/prds/proper-warm.md"])
            self.assertEqual(document.read_text(), "Proper Warm original evidence")
            self.assertEqual((root / "docs/prds/zava-forward.md").read_text(), "Zava Forward original evidence")
            self.assertEqual(log.read_text(), "[old evidence](../prds/proper-warm.md)")
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            again = rename.migrate(root, apply=False, logo=None)
            self.assertEqual(again["changes"], [])

    def test_logo_clears_the_existing_title_band(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(.68), Inches(.42), Inches(11.9), Inches(.72))
        title.text = "Proper Warm delivery challenge"
        logo = BytesIO()
        Image.new("RGB", (644, 225)).save(logo, "PNG")
        raw = BytesIO()
        presentation.save(raw)
        updated = Presentation(BytesIO(rename.migrate_pptx(raw.getvalue(), logo.getvalue())))
        picture = next(s for s in updated.slides[0].shapes if s.name == "Zava branding logo")
        self.assertLess(picture.top + picture.height, title.top)

    def test_infra_short_prefix_and_official_funding(self):
        self.assertEqual(
            rename.update_text("param workload = 'warm'\nid-warm-test sql-warm-test\nWarm Homes fund", "infra/a.bicep", {}),
            "param workload = 'forward'\nid-forward-test sql-forward-test\nWarm Homes fund",
        )

    def test_web_logo_and_styles(self):
        self.assertIn('src="/zava-logo.png"', rename.update_text(
            '<span class="logo" aria-hidden="true">PW</span>', "mvp/web/index.html", {},
        ))
        once = rename.update_text("", "mvp/web/style.css", {})
        self.assertEqual(rename.update_text(once, "mvp/web/style.css", {}), once)

    def test_cross_run_pptx_text_notes_and_metadata(self):
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = slide.shapes.add_textbox(0, 0, 1000000, 1000000).text_frame.paragraphs[0]
        paragraph.add_run().text = "Proper "
        paragraph.add_run().text = "Warm"
        slide.notes_slide.notes_text_frame.text = "Proper North Homes"
        presentation.core_properties.title = "Proper Warm"
        stream = BytesIO()
        presentation.save(stream)
        result = Presentation(BytesIO(rename.migrate_pptx(stream.getvalue(), None)))
        self.assertEqual(result.slides[0].shapes[0].text, "Zava Forward")
        self.assertIn("Zava North Homes", result.slides[0].notes_slide.notes_text_frame.text)
        self.assertEqual(result.core_properties.title, "Zava Forward")

    def test_logo_added_once(self):
        from pptx import Presentation
        from PIL import Image
        logo = BytesIO()
        Image.new("RGB", (644, 225), "white").save(logo, "PNG")
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        raw = BytesIO()
        presentation.save(raw)
        first = rename.migrate_pptx(raw.getvalue(), logo.getvalue())
        second = rename.migrate_pptx(first, logo.getvalue())
        result = Presentation(BytesIO(second))
        self.assertEqual(sum(s.name == "Zava branding logo" for s in result.slides[0].shapes), 1)

    def test_dry_run_then_apply_preserves_modes_and_untracked_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", root], check=True)
            file = root / "properwarm.py"
            file.write_text("print('Proper Warm')\n")
            file.chmod(0o755)
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            secret = root / ".env"
            secret.write_text("DO_NOT_TOUCH=Proper Warm")
            plan = rename.migrate(root, apply=False, logo=None)
            self.assertEqual(plan["changes"][0]["to"], "zavaforward.py")
            self.assertTrue(file.exists())
            result = rename.migrate(root, apply=True, logo=None)
            self.assertFalse(result["failures"])
            self.assertEqual((root / "zavaforward.py").read_text(), "print('Zava Forward')\n")
            self.assertEqual((root / "zavaforward.py").stat().st_mode & 0o777, 0o755)
            self.assertEqual(secret.read_text(), "DO_NOT_TOUCH=Proper Warm")

    def test_collision_stops_before_changes(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", root], check=True)
            for name in ("properwarm.py", "zavaforward.py"):
                (root / name).write_text("Proper Warm")
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            with self.assertRaisesRegex(ValueError, "collision"):
                rename.migrate(root, apply=True, logo=None)
            self.assertEqual((root / "properwarm.py").read_text(), "Proper Warm")

    def test_symlink_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", root], check=True)
            (root / "link.py").symlink_to("/etc/passwd")
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            with self.assertRaisesRegex(ValueError, "symlink"):
                rename.migrate(root, apply=True, logo=None)

    def test_bad_pptx_returns_partial_failure(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", root], check=True)
            (root / "broken.pptx").write_bytes(b"not a zip")
            subprocess.run(["git", "-C", root, "add", "."], check=True)
            result = rename.migrate(root, apply=False, logo=None)
            self.assertEqual(result["status"], "partial")
            self.assertEqual(result["failures"][0]["path"], "broken.pptx")

    def test_cli_apply_refuses_dirty_main_and_keeps_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            subprocess.run(["git", "init", "-q", "-b", "main", root], check=True)
            (root / "properwarm.py").write_text("Proper Warm")
            completed = subprocess.run(
                [sys.executable, str(Path(rename.__file__)), "--root", str(root),
                 "--apply", "--report", str(root.parent / "unused-zava-report.json")],
                capture_output=True, text=True,
            )
            self.assertEqual(completed.returncode, 2)
            self.assertIn("clean topic branch", completed.stderr)
            self.assertEqual((root / "properwarm.py").read_text(), "Proper Warm")


if __name__ == "__main__":
    unittest.main()
