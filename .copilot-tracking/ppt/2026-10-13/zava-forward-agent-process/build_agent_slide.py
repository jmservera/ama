from pathlib import Path
import argparse
import subprocess
import sys

from pptx import Presentation
from pptx.oxml.ns import qn
import yaml


def main():
    work = Path(__file__).resolve().parent
    root = work.parents[3]
    parser = argparse.ArgumentParser(description="Rebuild slide 13 with its YAML accessibility metadata")
    parser.add_argument("--builder", required=True, type=Path)
    parser.add_argument("--source", type=Path, default=root / "docs/capstone/deck/zava-forward-executive-draft.pptx")
    parser.add_argument("--output", type=Path, default=work / "slide-deck/zava-forward-executive-draft.pptx")
    args = parser.parse_args()
    if len(Presentation(args.source).slides) != 22:
        parser.error("Use the current 22-slide canonical deck, not the historical 21-slide source")
    content = yaml.safe_load((work / "content/slide-013/content.yaml").read_text())
    descriptions = {item["name"]: item["alt_text"] for item in content["elements"] if "alt_text" in item}
    if not descriptions or any(not value.strip() for value in descriptions.values()):
        parser.error("Slide YAML must supply nonempty image alt_text")
    subprocess.run([
        sys.executable, str(args.builder), "--content-dir", str(work / "content"),
        "--style", str(work / "content/global/style.yaml"), "--source", str(args.source),
        "--output", str(args.output), "--slides", "13",
    ], check=True)
    deck = Presentation(args.output)
    if len(deck.slides) != 22:
        raise ValueError("Partial rebuild changed the slide count")
    # The shared builder ignores alt_text; embed the YAML metadata before delivery.
    applied = set()
    for shape in deck.slides[12].shapes:
        if shape.name in descriptions:
            properties = shape._element.find(".//" + qn("p:cNvPr"))
            if properties is None:
                raise ValueError(f"Missing nonvisual properties for {shape.name}")
            properties.set("descr", descriptions[shape.name])
            applied.add(shape.name)
    if applied != set(descriptions):
        raise ValueError("Not every YAML alternative description matched a built shape")
    deck.save(args.output)


if __name__ == "__main__":
    main()
