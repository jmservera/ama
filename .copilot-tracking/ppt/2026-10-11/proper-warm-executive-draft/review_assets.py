"""Produce review assets from the actual final PPTX and PowerPoint exports."""
from pathlib import Path
import hashlib
import json
import re
from PIL import Image, ImageDraw
from pptx import Presentation
import yaml

root=Path(__file__).parent
deck=root/"slide-deck/proper-warm-executive-draft.pptx"
images=root/"slide-deck/validation/render-final-2"
prs=Presentation(deck)
files=sorted(images.glob("slide-*.jpg"))
assert len(files)==len(prs.slides)==17
assert [p.name for p in files]==[f"slide-{n:03d}.jpg" for n in range(1,18)]
canvas=Image.new("RGB",(1440,1680),"#dedede")
for j,f in enumerate(files):
    im=Image.open(f)
    assert im.size==(2000,1125)
    im.thumbnail((460,259))
    x=(j%3)*480+10
    y=(j//3)*280+18
    canvas.paste(im,(x,y))
    ImageDraw.Draw(canvas).text((x,y-15),f.stem,fill="black")
canvas.save(images/"contact-sheet.jpg")

def texts(shapes):
    for shape in shapes:
        if hasattr(shape,"shapes"):
            yield from texts(shape.shapes)
        elif shape.has_text_frame and shape.text:
            yield shape.text

lines=["<!-- markdownlint-disable-file -->","# Final PPTX text and notes — evidence review",
       "Extracted from the actual built PowerPoint, including native-group labels.",
       "14 main slides and 3 appendices. Main spoken notes: 1,991 words; citations excluded."]
spoken=0
for n,slide in enumerate(prs.slides,1):
    content=yaml.safe_load((root/f"content/slide-{n:03d}/content.yaml").read_text(encoding="utf-8"))
    notes=slide.notes_slide.notes_text_frame.text.strip()
    assert notes==content["speaker_notes"].strip(),f"Notes mismatch on slide {n}"
    body=notes.split("\n\nEvidence:")[0]
    if n<=14: spoken+=len(re.findall(r"\b[\w’'-]+\b",body))
    lines.extend([f"## Slide {n:02d} — {content['title'].replace(chr(10),' ')}",
                  "### Visible text","\n\n".join(texts(slide.shapes)),"### Speaker notes",notes])
assert spoken==1991
text="\n\n".join(lines)
assert not re.search(r"\blorem\b|\bipsum\b|\bxxxx\b|LINKED_IMAGE_NOT_EMBEDDED",text,re.I)
(root/"research/final-pptx-text-and-notes.md").write_text(text,encoding="utf-8")
manifest={"deck":str(deck),"sha256":hashlib.sha256(deck.read_bytes()).hexdigest(),
          "slide_count":len(prs.slides),"main_slides":14,"appendices":3,
          "main_spoken_note_words_excluding_citations":spoken,
          "image_directory":str(images),"image_count":len(files),
          "resolution":[2000,1125],"notes_match_yaml":True,"placeholder_scan":"clean"}
(root/"slide-deck/validation/final-artifact-manifest.json").write_text(
    json.dumps(manifest,indent=2),encoding="utf-8")
print(json.dumps(manifest,indent=2))
