"""Generate the bounded 24-slide presentation adaptation from the current baseline."""
from __future__ import annotations

import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
import shutil
import xml.etree.ElementTree as ET

from pptx import Presentation
import yaml

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]
BASE = ROOT.parent / "zava-forward-capstone-rebuild"
CONTENT = ROOT / "content"
CANONICAL = BASE / "slide-deck/zava-forward-executive-draft.pptx"
EXPECTED = "7b25fd7e77d1ed7d9aa908cf3da429bcfeae104beadea5ac90c1e97dbc71fbb4"
FOREST, INK, CREAM, MOSS, WARM = "#173F35", "#183A31", "#F7F2E7", "#D5E2D4", "#E5B96A"
TIMES = [35, 25, 35, 40, 25, 30, 60, 40, 30, 55, 40, 110, 60, 50, 30, 50, 45, 35]
TARGETS = list(range(1, 25))
TITLES = {
    1: "Back the evidence. Keep the resident promise.",
    3: "Each request needs rules, people and bounded AI.",
    4: "Compare routes; keep the exercise choice provisional.",
    10: "One API supports the work. People retain authority.",
    11: "Models propose. People and rules control change.",
    12: "Five real proposal paths. No delegated decisions.",
    17: "Earn the next gate; do not repeat the last proof.",
    18: "Endorse the next acceptance-and-value proposal.",
    19: "One runtime. Proposals are not domain approvals.",
    21: "Read the evidence label before the claim.",
    22: "Compare equivalent routes before selecting a winner.",
    23: "Operate the core without overstating assurance.",
    24: "Deferred services are not the five-assistant runtime.",
}
S16 = (
    "S16: docs/architecture/README.md, executive and technical Excalidraw/SVG/PNG views, "
    "docs/architecture/assets/README.md; 14 October 2026. Static source baseline "
    "062434a1fedbf6a87d1374c767c2bf0a11491ab6, not a fresh inventory or acceptance. "
    "Source Excalidraw remains editing authority; this PPTX is a simplified native adaptation. "
    "Individual unmodified official Azure SVG assets are locally rasterised for product icons; "
    "native containers, labels and connectors remain editable."
)
ARCH_ORACLE = (
    "Refresh documentation intake: f94f53a56c36d38ec36f7f662595912084571380. "
    "Architecture source: 062434a1fedbf6a87d1374c767c2bf0a11491ab6. "
    "Runtime source: cd14b3b379026a87ebfa3e9058ac10c0df7888eb. "
    "Historical cabe05c context below is not the current runtime or this documentation intake.\n"
    "Current five assistants are modules in ONE Python API Container App, not five microservices "
    "or the deferred queue-worker estate. Direct Azure OpenAI; no deployed Foundry Agent Service. "
    "SQL owns domain/proposal state, dedupe, reservations and fenced leases. Inference is outside "
    "database locks; the provider has no DB/workflow credentials. Accepting a proposal changes "
    "proposal/audit state, never resident, professional, work-release, quality or Finance approval.\n"
    "Normal five-assistant authority is absent; proactive OFF. This does not disable every legacy "
    "AI draft route. Ready = generated, unreviewed. Five operator paths, five calls, 6884 tokens, "
    "15 ledger events and full SQL content/restart proof are technical observations only, not "
    "staff-browser, human source, semantic or value acceptance. Restart is not restore acceptance.\n"
    "Development public-network exceptions coexist with configured private controls; PaaS is not "
    "VNet-hosted. Source-configured diagnostics are not full app tracing. External immutable audit "
    "digest anchoring is unimplemented. M365/import/queue workers are conditional or deferred. "
    "User.Read own-photo permission is not Sites.Selected plus a site grant. SharePoint topology "
    "parameter is not an implemented adapter. Front Door Premium/WAF is not deployed.\n"
    "Tokens: separate private Blob; daily 03:00 UTC renewal of a five-day user-delegation SAS; "
    "one vault secret write, distinct maintenance identity, no browser token retrieval. "
    "Budget 80% actual/100% forecast notifications are not a hard cap. Managed-identity ACR image "
    "pull/operator build-deploy is distinct from CI tests/scan/SBOM, not automatic Azure deployment.\n"
    "The owner explicitly reauthorised synthetic ledger recreation on 14 October; previous "
    "exercise audit/recommendations reset, not an unchanged migration. Retained 60,000-stock "
    "lineage; 980 cases, 38 booked visits and 11 private staff bindings are aggregate fixture "
    "counts only. No private identifiers are reproduced. Historical Decision 1 is not a current "
    "five-assistant runtime dependency. Finance fixtures and 20%/25% benefits are hypotheses, "
    "not measured success. No new legal currentness, expenditure, rollout or production approval."
)
BRIDGES = {
    3: "That tells us where AI may help, not which platform should win.",
    4: "Whichever route wins, the people and accountable services stay.",
    9: "Here is the small service core supporting that decision trace.",
    10: "The services explain where work lives; the next view explains who may change it.",
    11: "Now separate that design boundary from the evidence actually recorded.",
    12: "Those five proposal paths serve these five human-owned stages.",
    17: "The next decision is to prepare a scoped proposal, not enable a rollout.",
}


def text(name, value, x, y, w, h, size=20, bold=False, color=INK):
    return dict(type="textbox", name=name, text=value, left=x, top=y, width=w, height=h,
                font="Liberation Sans", font_size=size, font_color=color, font_bold=bold,
                vertical_anchor="top", auto_size="none", word_wrap=True,
                margin_left=.02, margin_right=.02, margin_top=.02, margin_bottom=.02,
                line_spacing=1.02, space_before=0, space_after=0)


def box(name, x, y, w, h, fill=MOSS, dashed=False):
    return dict(type="shape", shape="rounded_rectangle", name=name, left=x, top=y,
                width=w, height=h, fill=fill, line_color=FOREST, line_width=1.3,
                dash_style="dash" if dashed else "solid",
                description="Native editable service or person container; text follows in reading order.")


def flow(name, points, dashed=False, both=False):
    elements = []
    for i, (a, b) in enumerate(zip(points, points[1:]), 1):
        elements.append(dict(type="connector", name=f"{name}-segment-{i}",
                             connector_type="straight", begin_x=a[0], begin_y=a[1],
                             end_x=b[0], end_y=b[1], line_color=FOREST, line_width=1.8,
                             dash_style="dash" if dashed else "solid",
                             head_end="triangle" if both and i == 1 else "none",
                             tail_end="triangle" if i == len(points)-1 else "none"))
    return elements


def node(elements, name, title, body, x, y, w, h, *, human=False, dashed=False,
         title_size=20, body_size=18, title_h=.34, body_y=.48):
    body_y = max(body_y, title_h+.12)
    elements.extend([
        box(name+"-background", x, y, w, h, CREAM if human else MOSS, dashed),
        text(name+"-title", title, x+.12, y+.08, w-.24, title_h, title_size, True),
        text(name+"-body", body, x+.12, y+body_y, w-.24, h-body_y-.06, body_size),
    ])


def band(elements, name, value, y=6.35, h=.55, size=18):
    elements.extend([box(name+"-background", .55, y, 12.23, h, WARM),
                     text(name, value, .70, y+.08, 11.93, h-.12, size, True)])


def peers(elements, entries, y=2.02, h=3.72, title_h=.94, title_size=24,
          body_size=22, fill=MOSS, prefix="peer"):
    w = (12.23-.22*(len(entries)-1))/len(entries)
    for i, (title, body) in enumerate(entries, 1):
        x = .55+(i-1)*(w+.22)
        elements.extend([
            box(f"{prefix}-{i}-background", x, y, w, h, fill),
            text(f"{prefix}-{i}-title", title, x+.18, y+.20, w-.36, title_h, title_size, True),
            text(f"{prefix}-{i}-body", body, x+.18, y+.20+title_h+.17, w-.36,
                 h-title_h-.57, body_size),
        ])


def link(elements, name, label, target, x, y, w):
    e = text(name, label, x, y, w, .30, 16, True, FOREST)
    e["target_slide"] = target
    e["description"] = label
    elements.append(e)


def main10(e):
    # Six source relationships, each drawn as editable native connector segments.
    e.extend(flow("exec-staff-api", [(3.15, 3.18), (3.50, 3.18)]))
    e.extend(flow("exec-people-api", [(3.15, 4.65), (3.50, 4.65)]))
    e.extend(flow("exec-entra-api", [(6.0, 2.77), (6.0, 3.05)], True))
    e.extend(flow("exec-api-model", [(8.50, 3.90), (9.05, 3.90)], True, True))
    e.extend(flow("exec-api-sql", [(4.70, 4.85), (4.70, 5.25)]))
    e.extend(flow("exec-api-blob", [(7.30, 4.85), (7.30, 5.25)], True))
    node(e, "exec-staff", "Staff assist\nresidents", "HUMAN | requests\nPhone/letter/help.\nNo resident portal.",
         .55, 2.02, 2.60, 2.0, human=True, title_size=18, title_h=.65, body_y=.85)
    node(e, "exec-people", "People own\napproval", "HUMAN | approves\nResident + Finance.\nProfessionals; QA.",
         .55, 4.26, 2.60, 2.0, human=True, title_size=18, title_h=.65, body_y=.85)
    e.extend([box("exec-entra-background", 3.50, 2.02, 5.0, .75, dashed=True),
              text("exec-entra-title", "Microsoft Entra ID | CONFIGURED", 3.62, 2.08, 4.76, .34, 18, True),
              text("exec-entra-body", "Microsoft sign-in; role-limited access.", 3.62, 2.43, 4.76, .34, 18)])
    node(e, "exec-api", "Azure Container Apps",
         "CORE | one workbench + Python API.\nWorkflow + five assistant modules.",
         3.50, 3.05, 5.0, 1.80, title_size=22, title_h=.50, body_y=.65, body_size=20)
    node(e, "exec-model", "Azure OpenAI", "CONDITIONAL | draft only\nScoped request; usage back.\nNo DB/workflow credentials.",
         9.05, 3.05, 3.73, 1.80, dashed=True, title_size=22, title_h=.50, body_y=.65)
    e.extend([
        box("exec-sql-background", 3.50, 5.20, 2.40, 1.17),
        text("exec-sql-title", "Azure SQL", 3.62, 5.28, 2.16, .34, 18, True),
        text("exec-sql-body", "CORE | records\nProposals / audit", 3.62, 5.66, 2.16, .68, 18),
        box("exec-blob-background", 6.10, 5.20, 2.40, 1.17, dashed=True),
        text("exec-blob-title", "Azure Blob", 6.22, 5.28, 2.16, .34, 18, True),
        text("exec-blob-body", "CONFIGURED\nFiles; SQL hashes", 6.22, 5.66, 2.16, .68, 18),
        text("exec-sql-evidence", "SQL persistence:\nrecorded technical evidence", 9.05, 5.32, 3.73, .80, 18, True),
    ])
    band(e, "exec-boundary", "Five-assistant authority absent / proactive OFF.\nHuman and value acceptance outstanding.", 6.42, .78)
    link(e, "runtime-link", "Runtime detail: slide 19", 19, 9.05, 2.20, 3.73)


def runtime19(e):
    # Identity -> verified principal -> one runtime; stores/model never acquire authority.
    e.extend(flow("runtime-sign-in", [(4.28, 2.55), (4.80, 2.55)], True))
    e.extend(flow("runtime-issuer", [(8.53, 2.55), (9.05, 2.55)], True))
    e.extend(flow("runtime-principal", [(10.90, 3.16), (10.90, 3.34), (6.65, 3.34), (6.65, 3.55)], True))
    e.extend(flow("runtime-api-model", [(4.80, 4.28), (4.28, 4.28)], True, True))
    e.extend(flow("runtime-api-sql", [(8.53, 4.28), (9.05, 4.28)]))
    e.extend(flow("runtime-review-api", [(6.65, 5.32), (6.65, 4.98)]))
    e.extend(flow("runtime-api-blob", [(8.15, 4.98), (8.15, 5.20), (10.90, 5.20), (10.90, 5.32)], True))
    node(e, "runtime-browser", "Staff browser | HUMAN",
         "Entra sign-in; HTTPS ingress.\nPublic /; protected /app + APIs.",
         .55, 2.02, 3.73, 1.18, human=True, title_size=18, body_size=16, body_y=.42)
    node(e, "runtime-entra", "Microsoft Entra ID",
         "CONFIGURED | issuer + roles.\nIdentity -> EasyAuth.",
         4.80, 2.02, 3.73, 1.18, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "runtime-auth", "EasyAuth | CONFIGURED",
         "HTTPS -> verified principal.\n443 -> 8088; /healthz.",
         9.05, 2.02, 3.73, 1.18, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "runtime-model", "Azure OpenAI",
         "CONDITIONAL / OFF; draft only.\nUsage; no DB/workflow rights.\nNo Foundry Agent Service.",
         .55, 3.55, 3.73, 1.45, title_size=18, body_size=16, body_y=.42, dashed=True)
    node(e, "runtime-api", "Azure Container Apps",
         "CORE | one Python API.\nFive assistant modules in one API.\nCapture/reserve; validate/store.",
         4.80, 3.55, 3.73, 1.45, title_size=18, body_size=16, body_y=.42)
    node(e, "runtime-sql", "Azure SQL | CORE",
         "Domain/proposal state + audit.\nDedupe, reservations, fences.\nInference outside locks.",
         9.05, 3.55, 3.73, 1.45, title_size=18, body_size=16, body_y=.42)
    node(e, "runtime-review", "Human proposal review | SEPARATE REQUEST",
         "Verify sources. API rechecks permission, source/version/status.\nAccept proposal -> proposal/audit only, NEVER domain approval.",
         .55, 5.32, 7.98, 1.20, human=True, title_size=18, body_size=18, body_y=.40)
    node(e, "runtime-blob", "Azure Blob",
         "CONFIGURED | files + hashes.\nAI: allowed metadata only.",
         9.05, 5.32, 3.73, 1.20, title_size=18, body_size=16, body_y=.42, dashed=True)
    e.append(text("runtime-authority", "Five-assistant authority absent. Proactive OFF. Controlled approvals stay separate.",
                  .55, 6.57, 12.23, .32, 18, True))
    link(e, "operations-link", "Operations: slide 23", 23, .55, 6.97, 3.6)
    link(e, "deferred-link", "Deferred topology: slide 24", 24, 4.5, 6.97, 4.3)


def operations23(e):
    e.extend(flow("renewal-token-key", [(4.80, 2.70), (4.28, 2.70)], True))
    e.extend(flow("renewal-vault-write", [(8.53, 2.70), (9.05, 2.70)], True))
    node(e, "ops-tokens", "Azure Blob | auth tokens",
         "CONFIGURED | private store\nEasyAuth19 uses SAS.\nRenewal gets delegation key.",
         .55, 2.02, 3.73, 1.55, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "ops-renewal", "Azure Container Apps",
         "CONFIGURED | renewal job\n03:00 UTC; daily five-day SAS.\nOwn maintenance identity.",
         4.80, 2.02, 3.73, 1.55, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "ops-vault", "Azure Key Vault",
         "CONFIGURED | one secret\nRenewal writes; auth reads.\nNo browser token retrieval.",
         9.05, 2.02, 3.73, 1.55, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "ops-registry", "Azure Container Registry",
         "CONFIGURED | MI image pull\nOperator build + deploy.\nCI: tests / scan / SBOM only.",
         .55, 3.78, 3.73, 1.55, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "ops-logs", "Azure Log Analytics",
         "CONFIGURED | logs + audit\nApp Insights + SQL/API logs.\nFull tracing not established.",
         4.80, 3.78, 3.73, 1.55, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "ops-alerts", "Azure Monitor alerts",
         "CONFIGURED | email alerts\n80% actual / 100% forecast.\nNot a hard spending cap.",
         9.05, 3.78, 3.73, 1.55, title_size=18, body_size=18, body_y=.42, dashed=True)
    e.extend([
        text("ops-network", "SOURCE-CONFIGURED: VNet, subnets, NSGs, private endpoints, DNS, AMPLS.\nDevelopment also permits public access. PaaS services are NOT VNet-hosted.",
             .55, 5.45, 12.23, .64, 18),
        text("ops-support", "CONFIGURED: API19 -> Graph own-photo User.Read; browser -> tile provider (list alternative).",
             .55, 6.15, 12.23, .36, 18),
    ])
    band(e, "ops-failures", "Model fails: manual work. SQL fails: writes may block. Token renewal fails: act before expiry.",
         6.51, .45, 18)


def deferred24(e):
    e.append(text("deferred-frame-label", "DEFERRED ADR 0001 | placeholder contracts, NOT the five-assistant runtime",
                  .55, 1.89, 12.23, .34, 18, True))
    e.extend(flow("deferred-enqueue", [(4.28, 2.96), (4.80, 2.96)], True))
    e.extend(flow("deferred-keda", [(8.53, 2.96), (9.05, 2.96)], True))
    e.extend(flow("deferred-import-queue", [(6.65, 3.79), (6.65, 3.98)], True))
    e.extend(flow("deferred-snapshot", [(4.28, 4.57), (4.80, 4.57)], True))
    e.extend(flow("deferred-quarantine", [(8.53, 4.57), (9.05, 4.57)], True))
    node(e, "deferred-orchestrator", "Orchestrator | DEFERRED",
         "5-minute task placeholder.\nCheckpoints -> API19.\nNot an AI supervisor.",
         .55, 2.29, 3.73, 1.50, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "deferred-queues", "Azure Queue Storage",
         "DEFERRED | agent-tasks\nimport-jobs + poison queues.\nKEDA tasks / import events.",
         4.80, 2.29, 3.73, 1.50, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "deferred-agents", "Agent app | DEFERRED",
         "KEDA placeholder, not current AI.\nContext/proposals -> API19.\nModels19; no data-store tools.",
         9.05, 2.29, 3.73, 1.50, title_size=18, body_size=16, body_y=.42, dashed=True)
    node(e, "deferred-m365", "M365 | CONDITIONAL",
         "SharePoint/Excel snapshot.\nSites.Selected + site grant.\nOne-way; not photo rights.",
         .55, 3.98, 3.73, 1.50, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "deferred-import", "Azure Container Apps",
         "DEFERRED | import-jobs stub.\nLimited SQL; Blob quarantine.\nNo workbook write-back.",
         4.80, 3.98, 3.73, 1.50, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "deferred-containers", "Azure Blob | DEFERRED",
         "imports / quarantine\nexports / audit-archive\nCurrent API + SQL: see19.",
         9.05, 3.98, 3.73, 1.50, title_size=18, body_size=18, body_y=.42, dashed=True)
    node(e, "deferred-edge", "Azure Front Door Premium/WAF",
         "NOT DEPLOYED | intended edge; ingress direct.",
         .55, 5.56, 5.95, .88, title_size=18, body_size=18, body_y=.40, dashed=True)
    node(e, "deferred-anchor", "Audit digest | NOT IMPLEMENTED",
         "Unlocked policy is not immutable anchoring.",
         6.83, 5.56, 5.95, .88, title_size=18, body_size=18, body_y=.40, dashed=True)
    e.append(text("deferred-teams", "CONDITIONAL: ops23 -> Teams Workflows. Never approval, consent or safety response.",
                  .55, 6.54, 12.23, .35, 18, True))


ICON_NODES = {
    10: {"exec-entra": "entra", "exec-api": "aca", "exec-model": "openai",
         "exec-sql": "sql", "exec-blob": "blob"},
    19: {"runtime-entra": "entra", "runtime-model": "openai", "runtime-api": "aca",
         "runtime-sql": "sql", "runtime-blob": "blob"},
    23: {"ops-tokens": "blob", "ops-renewal": "aca", "ops-vault": "vault",
         "ops-registry": "registry", "ops-logs": "logs", "ops-alerts": "monitor"},
    24: {"deferred-queues": "queue", "deferred-import": "aca",
         "deferred-containers": "blob", "deferred-edge": "edge"},
}


def product_icons(slides):
    """Unmodified single-product assets only; preserve aspect and all native diagram objects."""
    import cairosvg
    from PIL import Image
    rows = []
    for number, nodes in ICON_NODES.items():
        elements = slides[number]["elements"]
        folder = CONTENT / f"slide-{number:03d}" / "images"
        for node_name, asset in nodes.items():
            source = REPO / f"docs/architecture/assets/azure/{asset}.svg"
            svg = ET.fromstring(source.read_bytes())
            viewbox = [float(v) for v in svg.attrib["viewBox"].split()]
            ratio = viewbox[2] / viewbox[3]
            png = folder / f"azure-{asset}.png"
            # Both raster dimensions derive from source aspect; neither crop nor recolour.
            pw, ph = round(240 * min(ratio, 1)), round(240 / max(ratio, 1))
            cairosvg.svg2png(bytestring=source.read_bytes(), write_to=str(png),
                            output_width=pw, output_height=ph)
            label = next(e for e in elements if e["name"] == node_name+"-title")
            x, y = label["left"], label["top"]
            label["left"] += .39
            label["width"] -= .39
            width, height = .29 * min(ratio, 1), .29 / max(ratio, 1)
            icon = dict(type="image", name=f"azure-icon-{asset}-{node_name}",
                        path=f"images/azure-{asset}.png", left=x, top=y+.005,
                        width=width, height=height,
                        description=f"Official Azure product icon; adjacent native label: {label['text']}.")
            elements.append(icon)
            row = dict(slide=number, node=node_name, object=icon["name"],
                       source=str(source.relative_to(REPO)),
                       source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                       png=str(png.relative_to(ROOT)),
                       png_sha256=hashlib.sha256(png.read_bytes()).hexdigest(),
                       source_viewbox=viewbox, source_ratio=ratio,
                       raster_pixels=list(Image.open(png).size), placement_inches=[width,height],
                       adjacent_label=label["text"], converter="CairoSVG "+cairosvg.__version__,
                       no_crop=True, no_recolour=True, source_unmodified=True,
                       provenance="docs/architecture/assets/README.md; local official Azure asset, not composite")
            png.with_suffix(".yaml").write_text(yaml.safe_dump(row, sort_keys=False))
            rows.append(row)
    (ROOT / "reports/azure-icon-provenance.json").write_text(json.dumps(rows, indent=2)+"\n")


def correct_map13(content):
    """Small uniform frame/spacing correction; no wording, font or horizontal geometry change."""
    for e in content["elements"]:
        if e["type"] != "textbox":
            continue
        e.update(line_spacing=1.0, space_before=0, space_after=0)
        if e["name"] == "slide-title":
            e["height"] = .57
        elif e["name"] == "source-footer":
            e.update(top=7.24, height=.25)
        elif e["name"] == "page-counter":
            e.update(text="13 / 24", top=7.24, height=.25)
        elif e["name"] == "map-status":
            e.update(top=1.24, height=.35)
        elif e["name"].endswith("-request-label"):
            e["height"] = .35
        elif e["name"].endswith("-request-body"):
            e.update(top=4.29, height=.58)
        elif e["name"].endswith("-proactive-label"):
            e.update(top=4.88, height=.35)
        elif e["name"].endswith("-proactive-body"):
            e.update(top=5.28, height=.62)
        elif e["name"] in ("map-funding-boundary", "map-hazard-boundary"):
            e["height"] = .33
            if e["name"] == "map-hazard-boundary":
                e["top"] = 6.90
    for e in content["elements"]:
        if e["name"].endswith("-capability-background"):
            e["height"] = 2.95


def base_elements(content, number):
    e = [deepcopy(a) for a in content["elements"] if a["name"] in
         ("deck-eyebrow", "brand-logo-backing", "Zava branding logo — only branding",
          "slide-title", "source-footer", "page-counter")]
    for a in e:
        if a["name"] == "source-footer":
            a.update(text="14 OCT 2026 | S16 + retained sources/status in notes", top=7.28, height=.21, font_size=10)
        if a["name"] == "page-counter":
            a.update(text=f"{number:02d} / 24", top=7.28, height=.20)
        if a["name"] == "deck-eyebrow":
            a["text"] = "ZAVA NORTH HOMES / ZAVA FORWARD / " + ("OPTIONAL REFERENCE" if number >= 19 else "EXECUTIVE PROPOSAL")
        if a["name"] == "slide-title":
            a["text"] = TITLES[number]
    return e


def source_mapping(slides):
    """Read the exact authoritative source IDs and map all cards/flows to actual objects."""
    aliases = {"people": "staff", "decision": "people", "identity": "entra", "workbench": "api",
               "sql": "sql", "blob": "blob", "model": "model"}
    exec_dest = {k: (10, "exec-"+v) for k, v in aliases.items()}
    tech_dest = {
        "staff": (19, "runtime-browser"), "entra": (19, "runtime-entra"), "maps": (23, "ops-support"),
        "graph": (23, "ops-support"), "identity": (23, "ops-registry"), "vault": (23, "ops-vault"),
        "model": (19, "runtime-model"), "auth": (19, "runtime-auth"), "api": (19, "runtime-api"),
        "sql": (19, "runtime-sql"), "blob": (19, "runtime-blob"), "tokens": (23, "ops-tokens"),
        "renewal": (23, "ops-renewal"), "registry": (23, "ops-registry"), "logs": (23, "ops-logs"),
        "insights": (23, "ops-logs"), "alerts": (23, "ops-alerts"), "agents": (24, "deferred-agents"),
        "queues": (24, "deferred-queues"), "orchestrator": (24, "deferred-orchestrator"),
        "import": (24, "deferred-import"), "other-containers": (24, "deferred-containers"),
        "edge": (24, "deferred-edge"), "m365": (24, "deferred-m365"),
    }
    tech_flows = {
        ("staff", "entra"): (19, "runtime-sign-in-segment-1"),
        ("staff", "auth"): (19, "runtime-browser-body"),
        ("entra", "auth"): (19, "runtime-issuer-segment-1"),
        ("auth", "api"): (19, "runtime-principal-segment-1"),
        ("api", "sql"): (19, "runtime-api-sql-segment-1"),
        ("api", "blob"): (19, "runtime-api-blob-segment-1"),
        ("api", "model"): (19, "runtime-api-model-segment-1"),
        ("vault", "auth"): (23, "ops-vault-body"),
        ("api", "graph"): (23, "ops-support"),
        ("staff", "maps"): (23, "ops-support"),
        ("renewal", "tokens"): (23, "renewal-token-key-segment-1"),
        ("renewal", "vault"): (23, "renewal-vault-write-segment-1"),
        ("auth", "tokens"): (23, "ops-tokens-body"),
        ("api", "registry"): (23, "ops-registry-body"),
        ("insights", "logs"): (23, "ops-logs-body"),
        ("sql", "logs"): (23, "ops-logs-body"),
        ("model", "alerts"): (23, "ops-alerts-body"),
        ("orchestrator", "queues"): (24, "deferred-enqueue-segment-1"),
        ("queues", "agents"): (24, "deferred-keda-segment-1"),
        ("queues", "import"): (24, "deferred-import-queue-segment-1"),
        ("import", "other-containers"): (24, "deferred-quarantine-segment-1"),
        ("agents", "api"): (24, "deferred-agents-body"),
        ("orchestrator", "api"): (24, "deferred-orchestrator-body"),
        ("agents", "model"): (24, "deferred-agents-body"),
        ("m365", "import"): (24, "deferred-snapshot-segment-1"),
        ("import", "sql"): (24, "deferred-import-body"),
        ("alerts", "m365"): (24, "deferred-teams"),
    }
    exec_flows = {
        ("staff", "api"): "exec-staff-api", ("entra", "api"): "exec-entra-api",
        ("api", "model"): "exec-api-model", ("api", "sql"): "exec-api-sql",
        ("api", "blob"): "exec-api-blob", ("people", "api"): "exec-people-api",
    }
    result = {}
    for view, dest in (("executive", exec_dest), ("technical", tech_dest)):
        source = REPO / f"docs/architecture/zava-forward-{view}-architecture.excalidraw"
        data = json.loads(source.read_text())
        nodes = [a for a in data["elements"] if a.get("customData", {}).get("role") == "service"]
        names = {a["id"]: a["customData"]["name"] for a in nodes}
        node_rows = []
        for a in nodes:
            name = a["customData"]["name"]
            n, prefix = dest[name]
            actual = [e["name"] for e in slides[n]["elements"] if e["name"].startswith(prefix)]
            assert actual, (view, name)
            node_rows.append(dict(source_id=a["id"], source_name=name, source_status=a["customData"]["status"],
                                  slide=n, native_objects=actual, sources=a["customData"].get("sources", [])))
        flows = []
        for a in data["elements"]:
            if a["type"] != "arrow":
                continue
            pair = (names[a["startBinding"]["elementId"]], names[a["endBinding"]["elementId"]])
            adapted_pair = tuple(aliases[v] for v in pair) if view == "executive" else pair
            n, obj = (10, exec_flows[adapted_pair]+"-segment-1") if view == "executive" else tech_flows[pair]
            assert any(e["name"] == obj for e in slides[n]["elements"]), (view, pair, obj)
            flows.append(dict(source_id=a["id"], source_from=pair[0], source_to=pair[1],
                              slide=n, native_object=obj,
                              representation="native connector" if "segment" in obj else "visible grouped interaction text"))
        result[view] = dict(source=str(source.relative_to(REPO)), source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                            nodes=node_rows, flows=flows)
    assert len(result["executive"]["nodes"]) == 7 and len(result["executive"]["flows"]) == 6
    assert len(result["technical"]["nodes"]) == 24 and len(result["technical"]["flows"]) == 27
    result["network_context"] = dict(slide=23, native_object="ops-network", not_a_deployed_service=True)
    return result


def generate():
    assert hashlib.sha256(CANONICAL.read_bytes()).hexdigest() == EXPECTED, "Canonical intake changed"
    for folder in ("content", "source", "reports", "slide-deck/validation"):
        (ROOT / folder).mkdir(parents=True, exist_ok=True)
    shutil.copytree(BASE / "content", CONTENT, dirs_exist_ok=True)
    slides = {i: yaml.safe_load((CONTENT / f"slide-{i:03d}/content.yaml").read_text()) for i in range(1, 23)}
    for i in (23, 24):
        slides[i] = deepcopy(slides[19])
        (CONTENT / f"slide-{i:03d}/images").mkdir(parents=True, exist_ok=True)
        shutil.copyfile(CONTENT / "slide-019/images/zava-logo.png", CONTENT / f"slide-{i:03d}/images/zava-logo.png")
        slides[i].update(slide=i, section="Optional appendix")
        slides[i]["speaker_notes"] = f"SLIDE {i:02d} - {TITLES[i]}\nMessage: {TITLES[i]}\nOptional appendix; no scheduled main-route time.\n"
    old = deepcopy(slides)
    byname = lambda n, name: next(e for e in slides[n]["elements"] if e["name"] == name)
    for i, c in slides.items():
        if i == 13:
            continue  # Original logical content and notes; bounded frame correction below.
        title = TITLES.get(i, c["title"])
        c["title"] = title
        byname(i, "slide-title")["text"] = title
        byname(i, "page-counter")["text"] = f"{i:02d} / 24"
        c["speaker_notes"] = re.sub(r"Planned allocation: \d+ seconds; not rehearsed.",
                                    f"Planned allocation: {TIMES[i-1]} seconds; not rehearsed." if i <= 18 else "Optional appendix, no main-route time.",
                                    c["speaker_notes"])
        c["speaker_notes"] = c["speaker_notes"].replace("1. Current documentation/source context is cabe05c;",
                                                                       "1. Historical documentation/source context was cabe05c;")
        c["speaker_notes"] = re.sub(r"^SLIDE \d+ [^\n]+", f"SLIDE {i:02d} - {title}", c["speaker_notes"], count=1)
        c["speaker_notes"] = "Message / evidence limit / planned script:\n" + c["speaker_notes"]
        c["speaker_notes"] += f"\n\nREFRESH BASELINE / UN SPOKEN REFERENCE:\n{ARCH_ORACLE}\n"
        if i in BRIDGES:
            c["speaker_notes"] += "\nSpoken transition: " + BRIDGES[i]
        if i >= 19:
            c["speaker_notes"] += "\nSpoken return: That is the current boundary; the decision remains which evidence gate to resolve next."
    byname(1, "decision-body").update(text="Endorse preparation\nof the next acceptance-\nand-value proposal.", font_size=24)
    byname(1, "boundary-text").update(text="Planning only - no new spending, rollout or production authority.", font_size=19)
    byname(5, "boundary-text").update(text="Excel could become reporting after a controlled transition.", font_size=22)
    # Replace4: preserve the complete original prerequisites in appropriate real notes.
    slides[3]["speaker_notes"] += "\nRetained slide4 prerequisite: Rank homes = approved policy plus sourced facts. Original example and notes:\n" + old[4]["speaker_notes"]
    slides[5]["speaker_notes"] += "\nRetained slide10 context, source authority and proposed integration (not delivered):\n" + old[10]["speaker_notes"] + "\n" + "\n".join(e.get("text","") for e in old[10]["elements"])
    slides[8]["speaker_notes"] += "\nSlot prerequisite: committed competent capacity plus resident measure-specific choice. A proposed slot is not consent."
    slides[14]["speaker_notes"] += "\nRetained prerequisite: accepted evidence plus separate application/award/claim/cash ledger states. A form is not funding. This is not technology price or budget."
    slides[4]["elements"] = base_elements(slides[4], 4)
    peers(slides[4]["elements"], [
        ("Improve workflow", "Least platform change.\nControls still to test."),
        ("Greenfield business platform", "Managed tools.\nNew adoption, licensing\nand support."),
        ("Azure-led hybrid", "Explicit contracts.\nBespoke engineering\nand operations."),
    ], title_size=24, body_size=22)
    band(slides[4]["elements"], "options-provisional", "Recommend hybrid for the bounded exercise only.\nProvisional baseline, NOT a customer investment winner.", 6.00, .80, 20)
    link(slides[4]["elements"], "options-link", "Equal-control appraisal: slide 22", 22, .55, 6.90, 6.0)
    byname(8, "boundary-text").update(text="A proposed slot is not consent. Silence/no-access is not refusal.", font_size=19)
    byname(11, "proposal-lane-label").update(text="PROPOSAL PATH / recorded operator integration; staff acceptance outstanding", font_size=18)
    byname(11, "domain-lane-label").update(text="SEPARATE CONTROLLED PATH / authority stays with people and rules", font_size=18)
    slides[11]["speaker_notes"] += "\nAccepting a proposal is a separate review request, not domain approval. Manual controlled work remains available; model calls do not span SQL locks."
    # Evidence12: immediate implication, unequal evidence/status treatment, no ambiguous OPEN.
    e = base_elements(slides[12], 12)
    e.append(text("proof-implication", "Integration and durability observed; usefulness not established.", .55, 1.95, 12.23, .40, 22, True))
    peers(e, [
        ("RECORDED TECHNICAL\nEVIDENCE", "5 distinct proposal paths\n5 real calls / 6,884 tokens\n15 ledger events\nControlled state unchanged.\nFull contents + usage survived restart."),
        ("NOT YET ESTABLISHED", "Staff browser/API + roles\nHuman source verification\nSemantic correctness + value\nReady = generated, unreviewed.\nAcceptance gates outstanding."),
    ], y=2.53, h=3.37, title_h=.70, title_size=20, body_size=20)
    next(a for a in e if a["name"] == "peer-2-background")["fill"] = CREAM
    band(e, "proof-authority", "Normal five-assistant authority absent. Proactive OFF.", 6.07, .48, 20)
    e.append(text("proof-provenance", "PUBLIC REPOSITORY EVIDENCE | 14 October | image05 / e0254cb3 | revision12",
                  .68, 6.67, 12.0, .33, 18, True))
    slides[12]["elements"] = e
    slides[12]["speaker_notes"] += "\nSynthetic context retained: 60,000 stock lineage; recreated 980 programme cases / 38 booked visits / 11 private staff bindings. Previous exercise audit/recommendations reset under explicit 14 October owner reauthorisation. No private details."
    byname(14, "fictional-label").update(text="FICTIONAL EXA-010/015 / not a technology price, award or budget")
    byname(17, "peer-2-title")["text"] = "2  Test"
    byname(17, "boundary-text").update(text="Outstanding gates: critical failure -> stop; no AI value -> simplify.", font_size=19)
    # Closing18: aligned owner/deliverable cards, prominent same planning-only ask.
    slides[18]["elements"] = base_elements(slides[18], 18)
    e = slides[18]["elements"]
    e.append(text("closing-ask", "Endorse preparation of the next acceptance-and-value proposal.",
                  .55, 1.98, 12.23, .74, 24, True, CREAM))
    peers(e, [
        ("CEO / Assets", "Owner: proposed sponsor\nDeliverable: accountable\nowners + scope."),
        ("CIO / Programme", "Owner: proposed operator\nDeliverable: human, access\nand recovery evidence."),
        ("CFO", "Owner: proposed reviewer\nDeliverable: whole-life cost\n+ scoped authority review."),
    ], y=2.98, h=2.70, title_h=.51, title_size=23, body_size=21, fill=CREAM)
    band(e, "closing-action", "Agree owners + return date (to agree). Continue, simplify or stop\nagainst pre-agreed tests.", 5.98, .82, 20)
    e.append(text("closing-boundary", "No new spending, procurement, rollout or production authority.",
                  .55, 6.87, 12.23, .36, 18, True, CREAM))
    slides[18]["speaker_notes"] += "\nCritical safety/rights/authority failure stops affected work. No incremental AI value -> simplify. Unresolved evidence -> return with named owner and status. Roles are proposals, not approved corporate delegations; date is to agree, not invented."
    # New architecture and technical progressive disclosure.
    for i, drawing in ((10, main10), (19, runtime19), (23, operations23), (24, deferred24)):
        slides[i]["elements"] = base_elements(slides[i], i)
        slides[i]["background"] = dict(fill=CREAM)
        drawing(slides[i]["elements"])
        slides[i]["speaker_notes"] += "\n\nS16 ARCHITECTURE / TEXT EQUIVALENT:\n" + S16 + "\n" + ARCH_ORACLE
        slides[i]["speaker_notes"] += "\nDiagram reading order: identity/request -> single API -> stores/model -> separate human authority. Connectors are editable; no composite architecture image. No native-viewer accessibility claim."
    # Retain old19 details where they still explain the runtime and operations limits.
    slides[19]["speaker_notes"] += "\nRETAINED TECHNICAL LIMITS FROM19:\n" + old[19]["speaker_notes"] + "\n" + "\n".join(e.get("text","") for e in old[19]["elements"])
    slides[23]["speaker_notes"] += "\nRETAINED OPERATIONS REFERENCE:\n" + old[19]["speaker_notes"] + "\nModel failure leaves manual work; SQL failure may block writes. Token renewal failure needs action before expiry. Safety intake has an independent named human repairs/hazard-service route, never AI/Teams dependent. SLO/load/unit cost, real sessions and full restore acceptance remain unmeasured. One deployable shares failure/scaling boundaries, a lean trade-off, not decoupled-microservice proof."
    slides[19]["speaker_notes"] += "\nThreats: prompt injection, source leakage, stale/duplicate proposals. Allowed sources and current role/source/version checks are separate from semantic verification. Fenced attempts/reservations constrain duplicates; a timeout does not mean zero usage; billing is not exactly once. Historical Decision 1 evaluation/autonomy gaps remain separate."
    slides[24]["speaker_notes"] += "\nSharePoint EVIDENCE_STORE topology parameter is not an implemented adapter. Agent KEDA 0-5, import event scale 0-2 and five-minute orchestrator are placeholder definitions, not current working code. Proposed workers have no SQL/Blob/Key Vault credential; API owns checkpoints/context/proposal contracts. Token maintenance is distinct from these workers. Additional evidence containers are defined; audit-archive policy is unlocked, complete immutable digest anchoring unimplemented."
    # Provenance and navigation changes preserve old source classification text in notes.
    byname(21, "source-footer")["text"] = "14 OCT 2026 | S01-S16 full source register + statuses in notes"
    byname(21, "peer-3-title")["text"] = "Recorded technical proof"
    slides[21]["speaker_notes"] += "\nSOURCE BASELINE LEGEND:\n" + S16 + "\n" + ARCH_ORACLE
    slides[22]["speaker_notes"] += "\nAzure hybrid is the provisional bounded exercise preference, not a customer investment winner. No existing Dataverse/Dynamics reuse. Test the same journey, safety hold, withdrawn choice, missing evidence and integration failure; include all transition/access/support/recovery/whole-life costs."
    byname(22, "boundary-text").update(text="Equivalent-scope appraisal needed. Exercise preference is NOT customer selection.", font_size=19)
    for i in range(19, 25):
        if i not in (19, 23, 24):
            byname(i, "source-footer").update(top=7.30, height=.17, font_size=9)
            byname(i, "page-counter").update(top=7.28, height=.20)
        link(slides[i]["elements"], "return-decision-link", "Return to decision: slide 18", 18, 9.05, 6.96, 2.95)
    link(slides[22]["elements"], "return-options-link", "Return to options: slide 4", 4, .55, 6.96, 4.5)
    link(slides[24]["elements"], "actual-runtime-link", "Actual runtime: slide 19", 19, .55, 6.96, 4.0)
    # Challenge navigation within existing reserve; never execution endpoints.
    slides[20]["speaker_notes"] += "\nQuestion routes: current runtime19; equivalent options22; operations23; deferred topology24. Listen/restate/answer/concede/assign unknown owner, then return18. No actual audience dialogue is claimed."
    byname(20, "boundary-text").update(text="Unknowns get an owner: runtime19 / options22 / operations23 / deferred24.", font_size=19)
    for e in slides[20]["elements"]:
        if e.get("text"):
            e["text"] = e["text"].replace("Human, semantic and value gates open",
                                        "Human, semantic and value acceptance outstanding")
    for i, target, label, x in ((18, 19, "Runtime detail:19", .55), (20, 22, "Options:22", .55)):
        link(slides[i]["elements"], f"challenge-{target}-link", label, target, x, 6.96 if i == 20 else 7.22, 3.0)
    # Do not let decision link collide with its footer; dark navigation gets cream text.
    e = byname(18, "challenge-19-link")
    e.update(text="Runtime: slide 19", top=.24, left=8.75, width=3.0, font_color=CREAM, font_size=16)
    byname(18, "deck-eyebrow")["width"] = 8.0
    byname(21, "assessment-open").update(text="Pitch, rehearsal + thought-leadership evidence outstanding.",
                                       width=8.20, height=.37)
    # Fixed-frame slack for measured PDF font metrics; no font shrinking or visible movement.
    # Map13 uses its own uniform spacing/frame correction.
    slack = {
        1: {"resident-scale": .08},
        4: {"options-provisional": .04},
        6: {"row-2-peer-1-body": .07, "row-2-peer-2-body": .07, "row-2-peer-3-body": .07},
        7: {"journey-hold-return": .06},
        8: {"human-hazard-route": .05},
        11: {"domain-peer-1-title": .07, "domain-peer-2-title": .07,
             "domain-peer-3-title": .07, "boundary-text": .07},
        15: {"peer-1-title": .10, "peer-2-title": .10,
             "peer-3-title": .10, "peer-4-title": .10},
        20: {"qa-row-1-peer-1-body": .10, "qa-row-1-peer-2-body": .10,
             "qa-row-2-peer-1-body": .08, "qa-row-2-peer-2-body": .08},
        22: {"option-peer-1-title": .11, "option-peer-2-title": .11, "option-peer-3-title": .11},
    }
    for i, entries in slack.items():
        for name, extra in entries.items():
            byname(i, name)["height"] += extra
    for i in range(1,25):
        byname(i,"page-counter").update(text=f"{i:02d} / 24", top=7.24, height=.25)
    correct_map13(slides[13])
    # Full expected-sequence glyph matching also exposed the retained slide15
    # three-line peer title's paragraph-after padding. Preserve words/fonts/frames.
    for e in slides[15]["elements"]:
        if e["name"].startswith("peer-") and e["name"].endswith("-title"):
            e.update(line_spacing=1.0, space_before=0, space_after=0)
    product_icons(slides)
    # External public evidence links are pinned source documents, never runtime execution.
    for i, doc in ((12, "infra/README.md#five-assistant-version-current-handover"), (21, "docs/architecture/README.md")):
        target = next(a for a in slides[i]["elements"] if a["name"] == ("proof-provenance" if i == 12 else "source-footer"))
        target["public_document"] = "https://github.com/jmservera/ama/blob/f94f53a56c36d38ec36f7f662595912084571380/" + doc
    mapping = source_mapping(slides)
    for view in ("executive", "technical"):
        for row in mapping[view]["flows"]:
            n = row["slide"]
            slides[n]["speaker_notes"] += f"\nSource flow {row['source_id']}: {row['source_from']} -> {row['source_to']}; native {row['native_object']}."
    for i in (10, 19, 23, 24):
        slides[i]["speaker_notes"] += "\nVISIBLE NATIVE DIAGRAM TEXT EQUIVALENT (unspoken):\n" + "\n".join(a.get("text", "") for a in slides[i]["elements"] if a["type"] == "textbox")
        slides[i]["speaker_notes"] += "\nFull source views: docs/architecture/zava-forward-executive-architecture.excalidraw / .svg / .png; docs/architecture/zava-forward-technical-architecture.excalidraw / .svg / .png. README and assets/README hold source/status/provenance."
    # Ground displaced-copy retention in the actual canonical package, not YAML alone.
    original_package = Presentation(CANONICAL)
    # Exact complete content, even map13 is retained as a read-only definition.
    for i, c in slides.items():
        if i != 13:
            if i <= 22:
                retained = [a.get("text", "") for a in old[i]["elements"]
                            if a.get("text") and a["name"] not in
                            ("deck-eyebrow","source-footer","page-counter")]
                for shape in original_package.slides[i-1].shapes:
                    if (shape.has_text_frame and shape.text.strip()
                            and shape.name not in ("deck-eyebrow","source-footer","page-counter")
                            and shape.text not in retained):
                        retained.append(shape.text)
                c["speaker_notes"] += (
                    "\n\nRETAINED BASELINE VISIBLE COPY / TRACEABILITY, UNSPOKEN:\n"
                    "Historical wording below is retained, not promoted to current fact or permission. "
                    "Current oracle/status and refreshed visible copy govern. In particular, old "
                    "'gates OPEN' meant outstanding acceptance, not permission; Excel reporting is a "
                    "conditional intended transition, not delivered integration.\n"
                    + "\n".join(retained)
                )
            c["speaker_notes"] = c["speaker_notes"].replace(
                "Current documentation/source context is cabe05c",
                "Historical documentation/source context was cabe05c")
        (CONTENT / f"slide-{i:03d}/content.yaml").write_text(yaml.safe_dump(c, sort_keys=False, allow_unicode=True, width=100))
    style = yaml.safe_load((CONTENT / "global/style.yaml").read_text())
    style["themes"][0]["slides"] += [23, 24]
    style["metadata"]["subject"] = "14 October 2026 architecture refresh; working candidate; planned not rehearsed"
    (CONTENT / "global/style.yaml").write_text(yaml.safe_dump(style, sort_keys=False, allow_unicode=True))
    p = Presentation(CANONICAL)
    assert len(p.slides) == 22
    for _ in range(2):
        s = p.slides.add_slide(p.slide_layouts[6])
        for shape in list(s.shapes):
            shape._element.getparent().remove(shape._element)
    p.save(ROOT / "source/structural-24-slides.pptx")
    (ROOT / "reports/source-to-native-mapping.json").write_text(json.dumps(mapping, indent=2)+"\n")
    manifest = []
    for i in range(1, 25):
        manifest.append(dict(slide=i, title=slides[i]["title"], new=i >= 23,
                             changed=i in TITLES or i in (5,8,13,14,15,20),
                             rebuilt=i in TARGETS, planned_seconds=TIMES[i-1] if i <= 18 else 0,
                             retention="Original13 logical content/native structure/fonts/notes retained; uniform spacing/frame correction and 13 / 24 metadata." if i == 13 else
                             "Original notes retained with updates; displaced content retained in associated reference notes/destinations."))
    (ROOT / "reports/slide-change-manifest.json").write_text(json.dumps(manifest, indent=2)+"\n")
    (ROOT / "reports/provenance.json").write_text(json.dumps(dict(
        canonical_source=str(CANONICAL.relative_to(REPO)), source_sha256=EXPECTED, source_slide_count=22,
        structural_source="source/structural-24-slides.pptx", output_slide_count=24,
        documentation_intake="f94f53a56c36d38ec36f7f662595912084571380",
        architecture_source="062434a1fedbf6a87d1374c767c2bf0a11491ab6",
        runtime_source="cd14b3b379026a87ebfa3e9058ac10c0df7888eb",
        planned_main_seconds=sum(TIMES), planned_qa_seconds=75, planned_contingency_seconds=30,
        evidence_walkthrough_seconds=110, fallback_within_walkthrough_max_seconds=15, rehearsed=False,
        font="Liberation Sans", added_service_icons=True,
        image_provenance="Retained branding plus unmodified individual official Azure assets; see azure-icon-provenance.json. No composite diagram.",
        targets=TARGETS, map13_preservation="Logical content, native structure, original notes, fonts and peer invariants retained; bounded spacing/frame/counter changes.",
    ), indent=2)+"\n")
    assert sum(TIMES) == 795


def finalise():
    """Apply only content-declared native slide links and alternative descriptions."""
    deck = ROOT / "slide-deck/zava-forward-executive-draft.pptx"
    prs = Presentation(deck)
    assert len(prs.slides) == 24
    for i, slide in enumerate(prs.slides, 1):
        c = yaml.safe_load((CONTENT / f"slide-{i:03d}/content.yaml").read_text())
        elements = {e["name"]: e for e in c["elements"]}
        for shape in slide.shapes:
            e = elements.get(shape.name, {})
            if "target_slide" in e:
                shape.click_action.target_slide = prs.slides[e["target_slide"]-1]
            if "public_document" in e:
                shape.click_action.hyperlink.address = e["public_document"]
            cnv = shape._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
            if cnv is not None and e.get("description"):
                cnv.set("descr", e["description"])
            if cnv is not None and shape.name.endswith("-background"):
                cnv.set("descr", "Container background; accompanying native text provides content.")
    prs.save(deck)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--finalise", action="store_true")
    args = parser.parse_args()
    finalise() if args.finalise else generate()
