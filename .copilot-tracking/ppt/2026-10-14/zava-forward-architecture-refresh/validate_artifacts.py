"""Validate the actual 24-slide package, local rendered words and authoritative mappings."""
from __future__ import annotations

from collections import Counter
from itertools import combinations
import csv
import hashlib
import importlib.metadata
import json
import posixpath
from pathlib import Path
import re
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import yaml

import baseline_validation as reused
from generate_content import ROOT, REPO, CANONICAL, EXPECTED, TIMES

REPORTS = ROOT / "reports"
DECK = ROOT / "slide-deck/zava-forward-executive-draft.pptx"
VALIDATION = ROOT / "slide-deck/validation"
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/package/2006/relationships"}


def out(name, data):
    (REPORTS / name).write_text(json.dumps(data, indent=2, ensure_ascii=False)+"\n")
    return data


def notes(prs, content):
    rows = []
    for i, s in enumerate(prs.slides, 1):
        rels = [r for r in s.part.rels.values() if r.reltype.endswith("/notesSlide")]
        actual = s.notes_slide.notes_text_frame.text
        assert len(rels) == 1
        rows.append(dict(slide=i, slide_part=str(s.part.partname), relationship_id=rels[0].rId,
                         relationship_target=rels[0].target_ref,
                         resolved_notes_part=str(rels[0].target_part.partname),
                         characters=len(actual), exact_yaml_match=actual == content[i]["speaker_notes"]))
    out("notes-association.json", rows)
    with (REPORTS / "notes-association-24-rows.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return rows


def integrity(prs):
    with ZipFile(CANONICAL) as a, ZipFile(DECK) as b:
        parts = [n for n in a.namelist() if n.startswith(("ppt/slideMasters/", "ppt/slideLayouts/", "ppt/theme/", "ppt/media/"))]
        compares = {n: a.read(n) == b.read(n) for n in parts}
        original = Presentation(CANONICAL)
        ids_a = [(s.slide_id, str(s.part.partname)) for s in original.slides]
        ids_b = [(s.slide_id, str(s.part.partname)) for s in prs.slides]
        map_part = str(original.slides[12].part.partname).lstrip("/")
        note_part = str(original.slides[12].notes_slide.part.partname).lstrip("/")
        map_rels = str(Path(map_part).parent / "_rels" / (Path(map_part).name+".rels"))
        unchanged_map = {n: a.read(n) == b.read(n) for n in (map_part, note_part, map_rels)}
        old_map, new_map = original.slides[12], prs.slides[12]
        old_shapes = {s.name: s for s in old_map.shapes}
        new_shapes = {s.name: s for s in new_map.shapes}
        logical = set(old_shapes) == set(new_shapes) and all(
            (s.shape_type == new_shapes[name].shape_type and
             (not s.has_text_frame or name == "page-counter" or s.text == new_shapes[name].text))
            for name,s in old_shapes.items())
        delta = []
        for name,s in old_shapes.items():
            t = new_shapes[name]
            changes = {k: [getattr(s,k),getattr(t,k)] for k in ("left","top","width","height")
                       if getattr(s,k) != getattr(t,k)}
            if s.has_text_frame:
                if s.text != t.text: changes["text"] = [s.text,t.text]
                for key in ("vertical_anchor","auto_size","margin_left","margin_right","margin_top","margin_bottom"):
                    assert getattr(s.text_frame,key) == getattr(t.text_frame,key), (name,key)
                assert [r.font.size for p in s.text_frame.paragraphs for r in p.runs] == [
                    r.font.size for p in t.text_frame.paragraphs for r in p.runs], name
                spacing = lambda sh: [(p.line_spacing,p.space_before,p.space_after) for p in sh.text_frame.paragraphs]
                if spacing(s) != spacing(t):
                    changes["paragraph_spacing"] = [spacing(s),spacing(t)]
            if changes: delta.append(dict(name=name,changes=changes))
        out("map13-layout-delta.json",dict(logical_content_retained=logical,
            original_notes_text_exact=old_map.notes_slide.notes_text_frame.text == new_map.notes_slide.notes_text_frame.text,
            fonts_insets_top_fixed_preserved=True, rows=delta))
        media_a = [n for n in a.namelist() if n.startswith("ppt/media/")]
        media_b = [n for n in b.namelist() if n.startswith("ppt/media/")]
        result = dict(
            working_path=str(DECK), sha256=hashlib.sha256(DECK.read_bytes()).hexdigest(),
            bytes=DECK.stat().st_size, canonical_sha256=hashlib.sha256(CANONICAL.read_bytes()).hexdigest(),
            canonical_untouched=hashlib.sha256(CANONICAL.read_bytes()).hexdigest() == EXPECTED,
            dimensions_identical=(prs.slide_width,prs.slide_height) == (original.slide_width,original.slide_height),
            dimensions_emu=[prs.slide_width,prs.slide_height],
            originals_1_22_identity_order_unchanged=ids_a == ids_b[:22],
            original_ids=ids_a, candidate_ids=ids_b,
            only_new_slides=[n+1 for n in range(22,len(prs.slides))],
            preserved_parts=compares, all_preserved=all(compares.values()),
            map13_byte_equivalence=unchanged_map, map13_unchanged=all(unchanged_map.values()),
            map13_logical_retained=logical,
            map13_notes_text_exact=old_map.notes_slide.notes_text_frame.text == new_map.notes_slide.notes_text_frame.text,
            retained_media=len(media_a), added_media=sorted(set(media_b)-set(media_a)),
        )
    return out("package-integrity.json", result)


def hyperlinks(prs, content):
    rows, errors = [], []
    for i, slide in enumerate(prs.slides, 1):
        declared = {e["name"]: e for e in content[i]["elements"]}
        for sh in slide.shapes:
            elem = declared.get(sh.name, {})
            if "target_slide" in elem:
                actual = sh.click_action.target_slide
                target = next((n for n,s in enumerate(prs.slides,1) if s == actual), None)
                rows.append(dict(slide=i, native_object=sh.name, expected=elem["target_slide"], actual=target,
                                 valid=target == elem["target_slide"]))
                if target != elem["target_slide"]:
                    errors.append([i, sh.name, "Wrong native hyperlink"])
            for rel in slide.part.rels.values():
                if rel.is_external:
                    url = rel.target_ref
                    if not (url.startswith("https://github.com/jmservera/ama/blob/") and "f94f53a" in url):
                        errors.append([i, sh.name, "Unexpected external relationship", url])
    required = {(10,19),(19,23),(19,24),(4,22)} | {(i,18) for i in range(19,25)}
    present = {(r["slide"],r["actual"]) for r in rows if r["valid"]}
    if not required <= present: errors.append(["Required links absent", list(required-present)])
    return out("hyperlink-validation.json", dict(pass_links=not errors, rows=rows, errors=errors,
                                                native_keyboard_viewer_testing="Not performed; XML target validity is not keyboard acceptance."))


def mappings(prs):
    m = json.loads((REPORTS / "source-to-native-mapping.json").read_text())
    checks = []
    for view in ("executive","technical"):
        d = json.loads((REPO / m[view]["source"]).read_text())
        src_nodes = {e["id"] for e in d["elements"] if e.get("customData",{}).get("role") == "service"}
        src_flows = {e["id"] for e in d["elements"] if e["type"] == "arrow"}
        assert src_nodes == {e["source_id"] for e in m[view]["nodes"]}
        assert src_flows == {e["source_id"] for e in m[view]["flows"]}
        for row in m[view]["nodes"] + m[view]["flows"]:
            names = {s.name for s in prs.slides[row["slide"]-1].shapes}
            objects = row.get("native_objects", [row.get("native_object")])
            checks.append(dict(view=view, source_id=row["source_id"], slide=row["slide"],
                               native_objects=objects, actual_objects_exist=all(n in names for n in objects)))
    counts = []
    for i in (10,19,23,24):
        s = prs.slides[i-1]
        pictures = [s for s in s.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        counts.append(dict(slide=i, shape_types=dict(Counter(str(s.shape_type) for s in s.shapes)),
                           pictures=[s.name for s in pictures],
                           product_icons=sum(s.name.startswith("azure-icon-") for s in pictures),
                           diagram_pictures=sum(not s.name.startswith("azure-icon-") and "branding" not in s.name for s in pictures),
                           connectors=sum(s.shape_type == MSO_SHAPE_TYPE.LINE for s in s.shapes)))
    return out("architecture-native-inventory.json", dict(
        all_mapped_to_actual_objects=all(c["actual_objects_exist"] for c in checks),
        executive_cards=7, executive_flows=6, technical_cards=24, technical_flows=27,
        source_coverage_verified=True, checks=checks, native_counts=counts,
        note="Some interactions are deliberately grouped visible native text, not an invented additional flow. Exact destinations in source-to-native-mapping.json."))


def linear(rgb):
    c = [int(rgb[i:i+2],16)/255 for i in (0,2,4)]
    return sum(k*(v/12.92 if v<=.04045 else ((v+.055)/1.055)**2.4) for k,v in zip((.2126,.7152,.0722),c))


def ratio(a,b):
    x,y=sorted([linear(a),linear(b)])
    return (y+.05)/(x+.05)


def contrast(prs):
    rows, errors = [], []
    for i, slide in enumerate(prs.slides,1):
        shapes = list(slide.shapes)
        bg = str(slide.background.fill.fore_color.rgb)
        for n,sh in enumerate(shapes):
            if not sh.has_text_frame or not sh.text.strip(): continue
            under = bg
            for other in shapes[:n]:
                if other.has_text_frame and not other.text.strip() and other.fill.type == 1:
                    if other.left <= sh.left and other.top <= sh.top and other.left+other.width >= sh.left+sh.width and other.top+other.height >= sh.top+sh.height:
                        under = str(other.fill.fore_color.rgb)
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    if not run.text.strip(): continue
                    fg=str(run.font.color.rgb)
                    r=ratio(fg,under)
                    rows.append(dict(slide=i,name=sh.name,foreground=fg,background=under,ratio=round(r,3),minimum=4.5,pass_contrast=r>=4.5))
                    if r<4.5: errors.append([i,sh.name,r])
    graphics=[]
    for i in (10,19,23,24):
        slide=prs.slides[i-1]
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.LINE:
                fg=str(sh.line.color.rgb)
                ratios={b:ratio(fg,b) for b in ("F7F2E7","D5E2D4","E5B96A")}
                graphics.append(dict(slide=i,name=sh.name,stroke=fg,against_adjacent_fills=ratios,pass_contrast=min(ratios.values())>=3))
    return out("measured-contrast.json",dict(pass_text=not errors,pass_connectors=all(r["pass_contrast"] for r in graphics),text=rows,connectors=graphics,errors=errors,
                                            method="Actual PPTX solid run colours and underlying contained shape/slide fill. Unmodified branding excluded; no transparency in meaningful diagram strokes."))


def connector_intersections(prs):
    findings=[]
    for i in (10,19,23,24):
        slide=prs.slides[i-1]
        frames=[s for s in slide.shapes if s.has_text_frame and s.text.strip()]
        for line in [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]:
            # All authored new segments are axis-aligned; ignore endpoint touches at card ports.
            x1,x2=sorted((line.begin_x,line.end_x)); y1,y2=sorted((line.begin_y,line.end_y))
            for sh in frames:
                left,right=sh.left,sh.left+sh.width
                top,bottom=sh.top,sh.top+sh.height
                hit=(x1==x2 and left<x1<right and max(y1,top)<min(y2,bottom)) or (y1==y2 and top<y1<bottom and max(x1,left)<min(x2,right))
                if hit: findings.append(dict(slide=i,connector=line.name,text_frame=sh.name))
    return out("connector-text-intersections.json",dict(pass_no_text_crossings=not findings,findings=findings))


def claims(prs):
    allnotes="\n".join(s.notes_slide.notes_text_frame.text for s in prs.slides)
    visible=lambda i:"\n".join(s.text for s in prs.slides[i-1].shapes if s.has_text_frame)
    checks={v:v.lower() in allnotes.lower() for v in (
        "cd14b3b379026a87ebfa3e9058ac10c0df7888eb","e0254cb3a0599c21802f0b6515f975550cf00ddce4849d04fd3fa3d50683af20",
        "ca-forward-test-api--0000012","f94f53a56c36d38ec36f7f662595912084571380","062434a1fedbf6a87d1374c767c2bf0a11491ab6",
        "ONE Python API","fenced leases","outside database locks","no DB/workflow credentials",
        "not domain approval","not five microservices","no deployed Foundry Agent Service","legacy",
        "6884","15 ledger","980","38","11 private","recreation","reset","03:00 UTC","five-day",
        "one vault secret","no browser token retrieval","not Sites.Selected","80% actual/100% forecast",
        "not a hard cap","not automatic Azure deployment","full app tracing","unimplemented",
        "SharePoint topology parameter","not an implemented adapter","not deployed","20%/25%",
        "Historical Decision 1","currency guarantee","RV-001","RV-002","110s","15s",
    )}
    proof=visible(12)
    checks.update({
        "ready_definition_visible":"Ready = generated, unreviewed" in proof,
        "acceptance_outstanding_visible":"Acceptance gates outstanding" in proof,
        "ambiguous_gates_open_removed":"gates OPEN" not in proof,
        "authority_off_visible":"authority absent" in proof and "Proactive OFF" in proof,
        "narrow_proof_implication_visible":"usefulness not established" in proof,
        "exact_receipt_counts_visible":all(t in proof for t in ("5 distinct","5 real","6,884","15 ledger")),
        "planning_only_same_ask":all("Endorse preparation" in visible(i) for i in (1,18)),
        "readable_no_authority_close":"No new spending, procurement, rollout or production authority" in visible(18),
        "no_private_details":"No private identifiers" in allnotes,
        "planned_time_exact":sum(TIMES)==795 and sum(TIMES)+75+30==900,
        "full_walkthrough_retained":all(t in prs.slides[11].notes_slide.notes_text_frame.text for t in ("0–15","15–40","40–70","70–95","95–110","pivot within")),
        "no_legacy_current_context": "Current documentation/source context is cabe05c" not in allnotes,
    })
    return out("claim-oracle-checks.json",dict(pass_claims=all(checks.values()),checks=checks,
                                             boundary="Local content/status regression, not independent service, legal or semantic acceptance."))


def frame_word_rows(prs, pdf, slides=None):
    """Supplement inherited token checks with actual PDF word rectangles near every native frame."""
    rows=[]
    with pymupdf.open(pdf) as d:
        assert len(d)==len(prs.slides)
        for i,(slide,page) in enumerate(zip(prs.slides,d),1):
            if slides is not None and i not in slides: continue
            sx=page.rect.width/prs.slide_width; sy=page.rect.height/prs.slide_height
            words=page.get_text("words")
            expanded = [(token,j) for j,w in enumerate(words) for token in reused.normalized(w[4]).split()]
            tokens = [t for t,j in expanded]
            for sh in slide.shapes:
                if not sh.has_text_frame or not sh.text.strip(): continue
                r=pymupdf.Rect(sh.left*sx,sh.top*sy,(sh.left+sh.width)*sx,(sh.top+sh.height)*sy)
                # Match the full expected token sequence first, even where a final line's
                # centre lies outside its frame. Repeated labels resolve to nearest native
                # origin, not a preservation allowlist or a clipped-region token proxy.
                expected = reused.normalized(sh.text).split()
                hits = [j for j in range(len(tokens)-len(expected)+1)
                        if tokens[j:j+len(expected)] == expected]
                selected = []
                if hits:
                    j = min(hits, key=lambda n: abs(words[expanded[n][1]][0]-r.x0)+abs(words[expanded[n][1]][1]-r.y0))
                    selected = [words[k] for k in sorted({k for t,k in expanded[j:j+len(expected)]})]
                outside=[dict(token=w[4],box=list(w[:4])) for w in selected
                         if w[0]<r.x0-1 or w[1]<r.y0-1 or w[2]>r.x1+1 or w[3]>r.y1+1]
                rows.append(dict(slide=i,name=sh.name,frame=list(r),words_matched_by_expected_sequence=len(selected),
                                 full_expected_sequence_matched=bool(hits),
                                 glyph_bounds_within_1pt=bool(hits) and not outside,overhang=outside))
    return rows


def bound_words(prs):
    rows = frame_word_rows(prs, VALIDATION/"slides.pdf")
    failures=[r for r in rows if not r["glyph_bounds_within_1pt"]]
    return out("rendered-word-bounds.json",dict(
        method="Full expected token sequence attributes actual PDF words before comparing complete word rectangles to literal native frame at 1pt tolerance. Includes out-of-frame word centres. Inherited clipped-region token checks run separately. No preservation exceptions; all 24 pages.",
        frames_checked=len(rows),pass_word_bounds=all(r["glyph_bounds_within_1pt"] for r in rows),
        failures=failures,checks=rows,dispositions=[],
        all_remaining_bounds_dispositioned=not failures))


def map_root_cause(prs):
    initial = REPORTS/"initial-candidate-receipts"
    baseline = REPORTS/"baseline-comparison.pdf"
    old_prs = Presentation(CANONICAL)
    old = frame_word_rows(old_prs, baseline, {13})
    first = frame_word_rows(Presentation(initial/"candidate.pptx"), initial/"slides.pdf", {13})
    final = frame_word_rows(prs, VALIDATION/"slides.pdf", {13})
    with pymupdf.open(baseline) as pdf:
        page = pdf[12]
        sx,sy = page.rect.width/old_prs.slide_width, page.rect.height/old_prs.slide_height
        token_rows = []
        for sh in old_prs.slides[12].shapes:
            if not sh.has_text_frame or not sh.text.strip(): continue
            rect = pymupdf.Rect(sh.left*sx-.5,sh.top*sy-.5,
                                (sh.left+sh.width)*sx+.5,(sh.top+sh.height)*sy+.5)
            missing = Counter(reused.normalized(sh.text).split())-Counter(reused.normalized(page.get_text("text",clip=rect)).split())
            token_rows.append(dict(name=sh.name,missing_tokens=dict(missing),pass_token_clip=not missing))
        metadata = pdf.metadata
    # Measure actual descendant-word/card breathing space, not frame-count proxy.
    map_slide = prs.slides[12]
    peers = []
    with pymupdf.open(VALIDATION/"slides.pdf") as pdf:
        sy = pdf[12].rect.height/prs.slide_height
        for n in range(1,6):
            frame = next(r for r in final if r["name"]==f"map-peer-{n}-proactive-body")
            sh = next(s for s in map_slide.shapes if s.name==frame["name"])
            expected = reused.normalized(sh.text).split()
            words = pdf[12].get_text("words")
            selected = [w for w in words if frame["frame"][0] <= (w[0]+w[2])/2 <= frame["frame"][2]
                        and frame["frame"][1] <= (w[1]+w[3])/2 <= frame["frame"][3]]
            card = next(s for s in map_slide.shapes if s.name==f"map-peer-{n}-capability-background")
            glyph_bottom = max(w[3] for w in selected)
            gap = (card.top+card.height)*sy-glyph_bottom
            peers.append(dict(peer=n,proactive_body_glyph_bottom_pt=glyph_bottom,
                              card_bottom_pt=(card.top+card.height)*sy,breathing_room_pt=gap,
                              pass_at_least_8pt=gap>=8))
    assert all(r["pass_at_least_8pt"] for r in peers)
    out("map13-glyph-root-cause.json",dict(
        baseline_export=str(baseline), baseline_pdf_metadata=metadata,
        initial_export=str(initial/"slides.pdf"),
        checker="Same complete expected-sequence PDF word attribution and literal frame 1pt tolerance in frame_word_rows for all three; includes words with out-of-frame centres.",
        explanation="The baseline 0.5pt helper checks expected tokens in an expanded region, not full glyph containment. "
                    "Same baseline/initial map13 produce the same conservative glyph rectangle failures. "
                    "They reveal genuinely undersized fixed frames under this exporter, compounded by 1.08 line spacing "
                    "and 5pt paragraph-after spacing on explicit line breaks. Not waived as metric false positives. "
                    "Uniform zero before/after and 1.0 line spacing, sufficient fixed frame heights, small peer-row "
                    "positions/card height delta and metadata correction resolve all final literal bounds. No font shrink.",
        baseline=old, initial_candidate=first, corrected_candidate=final,
        baseline_legacy_half_point_token_checks=token_rows,
        baseline_initial_same_literal_checks=old==first,
        corrected_peer_lower_card_breathing_room=peers,
        failure_counts=[sum(not r["glyph_bounds_within_1pt"] for r in rows) for rows in (old,first,final)]))


def icon_checks(prs):
    from PIL import Image
    declared = json.loads((REPORTS/"azure-icon-provenance.json").read_text())
    by_name = {(r["slide"],r["object"]):r for r in declared}
    rows, errors = [], []
    for i,slide in enumerate(prs.slides,1):
        for sh in slide.shapes:
            if sh.shape_type != MSO_SHAPE_TYPE.PICTURE or "branding" in sh.name: continue
            r = by_name.get((i,sh.name))
            if not r:
                errors.append([i,sh.name,"Undeclared/composite picture"]); continue
            png = ROOT/r["png"]
            source = REPO/r["source"]
            okay = (hashlib.sha256(sh.image.blob).hexdigest() == r["png_sha256"] ==
                    hashlib.sha256(png.read_bytes()).hexdigest() and
                    hashlib.sha256(source.read_bytes()).hexdigest() == r["source_sha256"] and
                    max(sh.crop_left,sh.crop_right,sh.crop_top,sh.crop_bottom) == 0 and
                    abs(sh.width/sh.height-r["source_ratio"]) < .0001)
            label = next(s for s in slide.shapes if s.name == r["node"]+"-title")
            okay = okay and label.text == r["adjacent_label"] and sh.left+sh.width < label.left
            # No product picture may intersect any text frame, even a different node/title.
            for t in slide.shapes:
                if t.has_text_frame and t.text.strip():
                    if min(sh.left+sh.width,t.left+t.width)>max(sh.left,t.left) and min(sh.top+sh.height,t.top+t.height)>max(sh.top,t.top):
                        errors.append([i,sh.name,t.name,"Picture/text collision"])
            rows.append(dict(slide=i,name=sh.name,sha_ratio_crop_adjacent_label_valid=okay))
            if not okay: errors.append([i,sh.name,"Invalid icon provenance/ratio/crop/label"])
    assert len(rows)==len(declared)
    return out("azure-icon-validation.json",dict(pass_icons=not errors,occurrences=len(rows),
        unique_assets=sorted({r["source"] for r in declared}),rows=rows,errors=errors,
        no_composite_diagrams=True,scope="Unmodified individual source icons only; native boxes/text/connectors preserved."))


def normalise_generated_markdown():
    trans=str.maketrans({"–":"-","—":"-","‘":"'","’":"'","“":'"',"”":'"',"≥":">=","≤":"<=","→":"->","≠":"!=","•":"*",
                         "£":"GBP", "ℹ":"INFO", "⚠":"WARNING", "✅":"PASS", "❌":"FAIL", "\ufe0f":""})
    for p in (VALIDATION/"deck-validation-report.md",VALIDATION/"shared-geometry-report.md"):
        if not p.exists(): continue
        t=p.read_text().translate(trans)
        if not t.startswith("---"):
            t=re.sub(r"^# (.+)$",r"## \1",t,flags=re.M)
            t=f"---\ntitle: {p.stem.replace('-',' ')}\nms.date: 2026-10-14\n---\n<!-- markdownlint-disable-file -->\n\n"+t
        t=re.sub(r"^(\s*)- ",r"\1* ",t,flags=re.M)
        p.write_text(t)


def main():
    p=Presentation(DECK)
    assert len(p.slides)==24
    content={i:yaml.safe_load((ROOT/f"content/slide-{i:03d}/content.yaml").read_text()) for i in range(1,25)}
    n=notes(p,content)
    pkg=integrity(p)
    g=out("native-peer-geometry.json",reused.geometry(p))
    rt=out("rendered-text-integrity.json",reused.rendered_text(p))
    wb=bound_words(p)
    map_root_cause(p)
    icons=icon_checks(p)
    icon_target = "/home/azureuser/.copilot/session-state/1e0d58ea-7ea9-4c65-93b7-66c89c481468/files/capstone-icon-tools"
    out("local-tools.json",dict(
        svg_converter_target=icon_target,requested_install="CairoSVG only; resolver dependencies isolated in same target",
        isolated_distributions=sorted([dict(name=d.metadata["Name"],version=d.version)
            for d in importlib.metadata.distributions(path=[icon_target])],key=lambda d:d["name"].lower()),
        actual_pymupdf_version=pymupdf.__version__,actual_pymupdf_path=pymupdf.__file__,
        uv_or_global_install=False,cloud_or_sdk_vision=False))
    arch=mappings(p)
    li=hyperlinks(p,content)
    co=contrast(p)
    ci=connector_intersections(p)
    cl=claims(p)
    sheets=reused.contact_sheets()
    raw=json.loads((VALIDATION/"shared-geometry-results.json").read_text())
    out("shared-warning-register.json",dict(status="Raw warnings retained. Final image-based dispositions recorded separately in handover/visual report.",
        total_findings=sum(len(s["issues"]) for s in raw["slides"]),slides=raw["slides"]))
    normalise_generated_markdown()
    flags=dict(slide_count=24,notes_24_exact=all(r["exact_yaml_match"] for r in n),
               package_intact=all(pkg[k] for k in ("canonical_untouched","dimensions_identical","originals_1_22_identity_order_unchanged","all_preserved","map13_logical_retained","map13_notes_text_exact")),
               geometry=g["pass_geometry"],tokens=rt["pass_rendered_text"],
               strict_all_frame_word_bounds=wb["pass_word_bounds"],icons=icons["pass_icons"],
               counters=all(next(s.text for s in slide.shapes if s.name=="page-counter")==f"{i:02d} / 24" for i,slide in enumerate(p.slides,1)),
               maps=arch["all_mapped_to_actual_objects"],links=li["pass_links"],text_contrast=co["pass_text"],
               connector_contrast=co["pass_connectors"],no_connector_text_crossing=ci["pass_no_text_crossings"],claims=cl["pass_claims"])
    out("validation-summary.json",dict(checks=flags,contact_sheets=sheets,working_sha256=pkg["sha256"],
                                      raw_strict_word_bounds_pass=wb["pass_word_bounds"],
                                      retained_map13_word_metric_findings=len(wb["failures"]),
                                      note="Fresh local artifact checks, not native viewer, rehearsal, panel or product acceptance."))
    print(json.dumps(flags,indent=2))
    for r in rt["failures"]:print("TOKEN",r["slide"],r["name"],r["missing_tokens"])
    for r in wb["failures"]:print("BOUNDS",r["slide"],r["name"],r["overhang"])
    for e in g["errors"]:print("GEOMETRY",e)
    for c,passed in cl["checks"].items():
        if not passed:print("CLAIM",c)
    if not all(flags.values()):raise SystemExit("Local artifact checks require fixes; inspect reports.")


if __name__=="__main__":main()
