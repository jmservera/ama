"""Record final local evidence and handover; never modify the presentation."""
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
import subprocess
import unicodedata

from pptx import Presentation

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]
REPORTS = ROOT / "reports"
DECK = ROOT / "slide-deck/zava-forward-executive-draft.pptx"


def read(name):
    return json.loads((REPORTS / name).read_text())


def ascii_text(text):
    trans = str.maketrans({
        "\u2013": "-", "\u2014": "-", "\u2018": "'", "\u2019": "'",
        "\u201c": '"', "\u201d": '"', "\u2265": ">=", "\u2264": "<=",
        "\u2192": "->", "\u2260": "!=", "\u2022": "*", "\u00a3": "GBP",
    })
    return unicodedata.normalize("NFKD", text.translate(trans)).encode("ascii", "ignore").decode()


def markdown(path, title, body):
    header = f"---\ntitle: {title}\nms.date: 2026-10-14\nstatus: checked-working-candidate\n---\n<!-- markdownlint-disable-file -->\n\n"
    body = re.sub(r"(?m)^\*(?=[0-9])", "* ", body)
    path.write_text(ascii_text(header + body).rstrip() + "\n")


def table(headers, rows):
    def cell(x):
        return str(x).replace("|", "/").replace("\n", " ")
    return "\n".join([
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
        *["| " + " | ".join(cell(x) for x in row) + " |" for row in rows],
    ]) + "\n"


def main():
    prs = Presentation(DECK)
    baseline = Presentation(REPO / "docs/capstone/deck/zava-forward-executive-draft.pptx")
    summary = read("validation-summary.json")
    assert hashlib.sha256(DECK.read_bytes()).hexdigest() == summary["working_sha256"]
    assert len(prs.slides) == 24
    assert all(summary["checks"].values())
    geom = read("native-peer-geometry.json")
    frames = {(f["slide"], f["name"]): f for f in geom["text_frames"]}
    groups = {
        10: [("exec-api", "exec-model"), ("exec-sql", "exec-blob")],
        19: [("runtime-browser", "runtime-entra", "runtime-auth"),
             ("runtime-model", "runtime-api", "runtime-sql"), ("runtime-review", "runtime-blob")],
        23: [("ops-tokens", "ops-renewal", "ops-vault"), ("ops-registry", "ops-logs", "ops-alerts")],
        24: [("deferred-orchestrator", "deferred-queues", "deferred-agents"),
             ("deferred-m365", "deferred-import", "deferred-containers"),
             ("deferred-edge", "deferred-anchor")],
    }
    peer_rows = []
    for slide, rows in groups.items():
        for names in rows:
            for role in ("title", "body"):
                group = [frames[slide, n + "-" + role] for n in names]
                keys = ("top_emu", "height_emu", "insets_emu")
                if role == "title":
                    keys += ("font_min_pt", "font_max_pt")
                signatures = [{k: f[k] for k in keys} for f in group]
                equal = all(s == signatures[0] for s in signatures)
                peer_rows.append(dict(slide=slide, names=names, role=role,
                                      zero_emu_and_font_drift=equal, signatures=signatures,
                                      body_font_sizes=[f["font_min_pt"] for f in group] if role == "body" else [],
                                      font_rule="Equal peer header fonts; body role-specific 16pt map labels / 18pt normal / 20pt central core retained, not claimed equal."))
    assert all(r["zero_emu_and_font_drift"] for r in peer_rows)
    (REPORTS / "architecture-peer-rows.json").write_text(json.dumps(peer_rows, indent=2) + "\n")

    def norm(text):
        return re.sub(r"\s+", " ", text).strip()
    retained = []
    excluded = {"deck-eyebrow", "source-footer", "page-counter"}
    for slide, old in enumerate(baseline.slides, 1):
        new = prs.slides[slide - 1]
        visible = norm("\n".join(s.text for s in new.shapes if s.has_text_frame))
        notes = norm(new.notes_slide.notes_text_frame.text)
        for shape in old.shapes:
            if not shape.has_text_frame or not shape.text.strip() or shape.name in excluded:
                continue
            text = norm(shape.text)
            where = "same-slide visible" if text in visible else (
                "same-slide relationship-bound reference notes" if text in notes else "MISSING")
            retained.append(dict(slide=slide, original_object=shape.name, original_text=shape.text,
                                 destination=where, retained=where != "MISSING"))
    assert all(r["retained"] for r in retained)
    (REPORTS / "original-content-retention.json").write_text(json.dumps(dict(
        excluded_metadata=sorted(excluded), all_unique_original_content_retained=True,
        original_frames_checked=len(retained), rows=retained), indent=2) + "\n")
    fonts = Counter(run.font.name for slide in prs.slides for sh in slide.shapes if sh.has_text_frame
                    for p in sh.text_frame.paragraphs for run in p.runs if run.text.strip())
    assert set(fonts) == {"Liberation Sans"}
    (REPORTS / "actual-font-inventory.json").write_text(json.dumps(dict(
        run_font_counts=dict(fonts), all_explicit_liberation_sans=True,
        nonessential_small_text_objects=["deck-eyebrow", "source-footer", "page-counter"],
        diagram_label_minimum_pt=16, normal_body_minimum_pt=18), indent=2) + "\n")

    raw = read("shared-warning-register.json")
    dispositions = []
    for s in raw["slides"]:
        for issue in s["issues"]:
            kind = issue["check_type"]
            if kind == "edge_margin":
                rationale = (
                    "Retained branding/provenance/counter or deliberately lower navigation/authority "
                    "strip inside the canvas. Inspected final full-size image and contact sheet; "
                    "no visible edge clipping or footer collision. Essential authority text is "
                    "body-sized, not a source-only disclaimer. Target-viewer/projection remains untested."
                )
            elif kind == "adjacent_gap":
                rationale = (
                    "Compact intentional title/card, lane/band or navigation gap below the generic "
                    "0.3-inch recommendation. Final image remains legible; actual native text frames "
                    "do not intersect, tokens are present and changed-frame PDF word bounds pass. "
                    "All24 strict full-sequence glyph bounds pass; no map13 exceptions."
                )
            else:
                assert kind == "title_clearance"
                rationale = (
                    "Generic title/body recommendation, including comparisons across neighbouring "
                    "peer columns. Intended peers have fixed equal starts and no text intersections. "
                    "Final image inspected; no title/body collision. Compact map13 has bounded uniform spacing/frame correction."
                )
            dispositions.append(dict(slide=s["slide_number"], **issue,
                                     disposition="retained-with-explicit-artifact-disposition",
                                     rationale=rationale))
    counts = Counter((r["severity"], r["check_type"]) for r in dispositions)
    assert len(dispositions) == raw["total_findings"]
    severities = Counter(r["severity"] for r in dispositions)
    kinds = Counter(r["check_type"] for r in dispositions)
    assert severities.get("error",0) == 0
    (REPORTS / "warning-dispositions.json").write_text(json.dumps(dict(
        raw_total=len(dispositions), errors=0, warning_count=severities["warning"], info_count=severities["info"],
        visual_review="All four contact sheets and full-size changed10/13/15/19/20/23/24 plus neighbours9/11/12/14/16/18/21/22 inspected locally.",
        no_blanket_clean_claim=True, rows=dispositions), indent=2) + "\n")

    visual_rows = []
    for i in range(1, 25):
        info = "No visible clipping or unintended text/connector collision observed in final local render."
        if i == 13:
            info = ("Dense semantic-original native map. All five capabilities, owners, OFF triggers, "
                    "funding/manual/hazard boundaries readable. Correct13 / 24 counter; uniform spacing, "
                    "slightly taller frames/cards and small common row shifts. Strict complete-word bounds pass.")
        elif i in (10, 19, 23, 24):
            info += " Native diagram inspected full-size with neighbours; explicit status labels and routed gutters."
        elif i in (4, 11, 12, 18, 21, 22):
            info += " Compact lower bands/navigation readable; generic gap/edge warnings retained."
        visual_rows.append([i, f"slide-deck/validation/slide-{i:03d}.jpg", info])
    markdown(REPORTS / "visual-inspection.md", "Final local rendered-asset inspection",
             "## Scope and limits\n\n* All four fresh contact sheets cover24 slides. Full-size150dpi "
             "JPGs inspected for changed10/13/15/19/20/23/24 plus neighbours9/11/12/14/16/18/21/22.\n"
             "* Local actual-asset expert inspection, not SDK model vision, "
             "native PowerPoint, human audience testing or rehearsal.\n* Initial overflow/wrapping "
             "defects were fixed and the full candidate was rerendered before this final inspection.\n\n" +
             table(["Slide", "Final image", "Finding / disposition"], visual_rows))

    manifest = read("slide-change-manifest.json")
    mapping = read("source-to-native-mapping.json")
    pkg = read("package-integrity.json")
    contrast = read("measured-contrast.json")
    inventory = read("architecture-native-inventory.json")
    content = f"""## Checked working candidate

* Status: complete within working-candidate boundary; not canonical publication or presentation/product/panel acceptance.
* Output: `{DECK}`.
* SHA256: `{pkg['sha256']}`; {pkg['bytes']} bytes; exactly 24 slides, main 1-18 and optional 19-24.
* Only 23/24 are new. Original 1-22 slide identities/order retained.
* Canvas: 12,191,695 x 6,858,000 EMU, identical to intake.
* Canonical intake remains SHA256 `{pkg['canonical_sha256']}`.
* Source masters, layouts, themes and retained media byte-identical; original slide identities and notes associations retained. Legitimate individual icon PNG media added.
* Native map13 wording (except counter), fonts, insets, TOP/fixed state, native structure and original relationship-bound notes text retained. Small uniform frame/spacing/layout corrections and `13 / 24` counter authorised by parent. Exact delta: `reports/map13-layout-delta.json`.
* No cloud/auth/model/live-app/private-record action, uv, agents/workflows, staging/commit/push or publication occurred. Only parent-authorised isolated CairoSVG pip target installed for demonstrated missing converter, with dependencies; no global environment changes.

## Inputs and provenance

* Documentation intake: f94f53a56c36d38ec36f7f662595912084571380; static architecture: 062434a1fedbf6a87d1374c767c2bf0a11491ab6; runtime: cd14b3b379026a87ebfa3e9058ac10c0df7888eb. These are different baselines.
* Runtime image05 digest: e0254cb3a0599c21802f0b6515f975550cf00ddce4849d04fd3fa3d50683af20; revision ca-forward-test-api--0000012.
* Grounding: complete original case/addendum, reference index/memory, company profile, housing-association primer and complete capstone assessment guide; approved UX brief read in full.
* Current 22-slide YAML/style/logo/helper patterns came from sibling `zava-forward-capstone-rebuild`, not historical 21-slide or 13 October builds.
* Native Excalidraw remains diagram editing authority. This PPTX is a simplified native presentation adaptation, not replacement source.
* Source views and assets under `docs/architecture/` are read-only. S01-S15 remain stable; S16 adds architecture.
* Images are retained branding plus20 placements of11 individual unmodified official Azure product assets across10/19/23/24. No composite diagram, cropped source fragment or whole architecture image. Correct aspect, colour, no crop and SVG/PNG SHA256 provenance: `reports/azure-icon-provenance.json`; actual package checks: `reports/azure-icon-validation.json`. No product icon on human/custom assistant nodes.
* Native source/object inventory: `reports/architecture-native-inventory.json`; source hashes and full mapping: `reports/source-to-native-mapping.json`.

## Slide change and content retention

* Every original nonmetadata visible text frame is retained in current visible text or its same-slide, relationship-bound historical reference notes: {len(retained)} frames checked in `reports/original-content-retention.json`. This is per-frame retention, not just a slide-count proxy.
* Original speaker/reference content retained with explicit current-versus-historical updates; historical ambiguous OPEN wording is not a current authority claim. Speaking script and unspoken reference content are distinct.
* Original4 rank/policy example also retained in3 notes; slot/capacity in8 and notes; funding/ledger in14 and notes. Old10 context also retained in5/10 notes; old19 State/Trust/Operate in19/23 notes.
* No original13 capability/owner/control/purpose wording, notes, font or role inventory changed. Uniform paragraph spacing and sufficient fixed frames, small common row shifts and card height delta resolve actual strict glyph failures; all24 counters now `NN / 24`.
* Exact map13 correction in inches: title height0.50->0.57; status top1.19->1.24/height0.30->0.35; each request-label height0.29->0.35; request-body top4.23->4.29/height0.57->0.58; proactive-label top4.86->4.88/height0.29->0.35; proactive-body top5.19->5.28/height0.59->0.62; capability-card height2.84->2.95; funding boundary height0.28->0.33; hazard top6.86->6.90/height0.27->0.33; footer/counter top7.24/height0.25. All text paragraph line spacing1.08->1.0, before0 and after5pt->0. No horizontal positions/widths/fonts/insets or four native arrows changed. Every lower body glyph has8.97pt card-bottom breathing room.
* Slide20 says human, semantic and value acceptance outstanding;10 uses a two-line18pt human/value acceptance status band;19 labels Azure OpenAI explicitly. Current proof/authority/deferred/network/audit boundaries retained.
* Stronger full-sequence glyph matching exposed one pre-existing15 peer title overflow that position-only matching missed. All four peer headers now use equal zero paragraph-after/1.0 spacing and0.10in additional fixed height, with words/fonts/body starts unchanged.

"""
    content += table(["Slide", "Current title", "Change", "Planned seconds", "Retention"], [
        [r["slide"], r["title"], "new" if r["new"] else (
            "targeted UX/content" if r["changed"] else "preserved content; metadata/timing only" if r["rebuilt"] else "byte-identical"),
         r["planned_seconds"], r["retention"]] for r in manifest])
    content += "\n## Approved UX findings: implementation disposition\n\n"
    content += table(["Rank", "Implemented change", "Evidence / remaining gate"], [
        [1, "Native executive10, runtime19, operations23 and deferred24 replace generic context/dense source image.", "7/24 source nodes and 6/27 flows verified against actual objects."],
        [2, "12 separates recorded evidence from not established; Ready generated/unreviewed; acceptance outstanding; authority absent/proactive OFF.", "Actual visible text and notes oracle pass; no staff/value acceptance implied."],
        [3, "1/18 repeat same planning-only ask;18 aligned proposed owner/deliverable cards and return date to agree.", "No spend/rollout/production boundary visible at body size; sponsor decision unobserved."],
        [4, "4 names all three alternatives and provisional exercise hybrid preference;22 fair controls, no scored winner.", "Prerequisite content redistributed and retained; real option appraisal still needed."],
        [5, "10 service core,11 proposal/control lanes,13 human-owned stages have distinct purpose/status.", "Single API not five services; no model-to-domain approval path."],
        [6, "12 integration/durability implication dominates;19 SQL leases/outside-lock inference;23 failure consequences.", "Five calls not semantic/value/restore proof; receipt remains narrowly scoped."],
        [7, "Retained palette and good slide grammar; new architecture adds native paths, not decorative charts.", "Warm caution/decision bands labelled; no fabricated ROI."],
        [8, "11/19 titles shortened after actual wrap checks;21 source pointer shortened; all frames TOP/fixed.", "Liberation Sans; native labels >=16, normal body >=18; map13 preserved."],
        [9, "Replaced main architecture instead of extending main route; exact timing retained.", "795+75+30=900 planned only; human rehearsal outstanding."],
    ])
    content += "\n## Full source-to-native mapping\n\n"
    for view in ("executive", "technical"):
        data = mapping[view]
        content += f"\n### {view.capitalize()} nodes\n\n* Source: `{data['source']}`; SHA256 `{data['source_sha256']}`.\n\n"
        content += table(["Source ID", "Source card", "Source state", "Slide", "Actual native objects"], [
            [r["source_id"], r["source_name"], r["source_status"], r["slide"], ", ".join(r["native_objects"])]
            for r in data["nodes"]])
        content += f"\n### {view.capitalize()} flows\n\n"
        content += table(["Source flow ID", "From -> to", "Slide", "Actual native object", "Representation"], [
            [r["source_id"], r["source_from"] + " -> " + r["source_to"], r["slide"],
             r["native_object"], r["representation"]] for r in data["flows"]])
    content += """
* Network context is23 `ops-network`, not an extra deployed service.
* Technical grouping: identity+registry -> ops-registry; Log Analytics+Application Insights -> ops-logs; Graph+tile provider -> ops-support. All names/interaction equivalents remain visible native text.
* Some source flows intentionally use grouped native text rather than an additional connector. No model-to-SQL authority is invented; current/deferred paths stay separate.

## Actual-output validation and warning dispositions

"""
    content += table(["Check", "Actual result / report"], [
        ["Package preservation", "24 count, original identities/canvas/masters/layouts/themes/media retained; map13 logical content/native structure/fonts/notes retained. reports/package-integrity.json; map13-layout-delta.json"],
        ["Notes", "24/24 actual relationship-bound notes exactly equal YAML; notes13 resolves to notesSlide22, not ordinal13. reports/notes-association.json and CSV"],
        ["Native diagrams", "All mapped native objects exist; no composite diagram pictures; retained logo plus20 individual Azure product icons. reports/architecture-native-inventory.json; azure-icon-validation.json"],
        ["Frames and peers", "Every nonempty frame TOP/fixed; zero native text-frame pair intersections; reused peer checks pass. reports/native-peer-geometry.json"],
        ["Architecture peers", f"{len(peer_rows)} intended title/body peer groups have zero EMU/inset drift; peer header fonts equal. Role-specific body label sizes explicitly retained. reports/architecture-peer-rows.json"],
        ["Fonts", "Every nonempty run explicitly Liberation Sans. Metadata-only 10/11pt exceptions enumerated. reports/actual-font-inventory.json"],
        ["Render", "Fresh LibreOffice full24-page PDF and all24 JPGs at150dpi; four contact sheets. slide-deck/validation/"],
        ["Rendered content", "Actual per-frame expected tokens pass on all24 pages; 387 native frames traversed for glyph bounds. reports/rendered-text-integrity.json; rendered-word-bounds.json"],
        ["Strict word bounds", "TRUE: complete expected-token sequences attribute full PDF word rectangles, including out-of-frame centres, to literal native frames at1pt. All387 frames pass, no exceptions. Baseline/initial13 both24 findings; corrected13 zero. reports/map13-glyph-root-cause.json"],
        ["Shared properties", "0 issues across24 slides. slide-deck/validation/deck-validation-results.json and report"],
        ["Shared geometry", f"{len(dispositions)} retained: {kinds['edge_margin']} edge-margin, {kinds['adjacent_gap']} adjacent-gap, {kinds['title_clearance']} title-clearance; {severities['warning']} warnings / {severities['info']} info / zero errors. Raw shared-geometry-results.json"],
        ["Warnings", f"All{len(dispositions)} rows retain original severity/location/description plus explicit class-based disposition. reports/warning-dispositions.json; final visual-inspection.md"],
        ["Contrast", f"All actual substantive text >=4.5:1; meaningful new connectors >=3:1. Measured {len(contrast['text'])} text runs and {len(contrast['connectors'])} connectors. reports/measured-contrast.json"],
        ["Routes", "No new architecture connector intersects a native text frame. reports/connector-text-intersections.json"],
        ["Links", "Actual slide relationships valid:10->19,19->23/24,4->22,all19-24->18;22->4 retained. Public sources pinned document URLs only. reports/hyperlink-validation.json"],
        ["Claims", "Full current/historical/status oracle and exact timing pass; reports/claim-oracle-checks.json. Presence/status checks supplement, not replace, source/visual review."],
    ])
    content += """
### Retained warnings are not an all-clear

* Edge findings: common branding/provenance/counter, product icons and lower navigation/authority strips. All lie inside canvas and were viewed. Small provenance is nonessential, with complete notes. Projection legibility still needs a human check.
* Gap findings: compact title/card/lane/band/navigation gaps below generic0.3-inch guidance. Final actual images, full expected tokens and complete word bounds show no collision; no map13 exemption.
* Title info: compact fixed headers and cross-column generic comparisons. Intended peer starts/heights/insets/fonts measured equal; no native text intersections. Not a claim of generic validator clearance.
* Map13: baseline and initial candidate both24 same-checker failures, corrected candidate zero. Old0.5pt expected-token clip checks were not glyph containment. Undersized fixed frames/paragraph-after spacing were corrected, not waived as font metrics. Initial raw failed receipts and PDF/candidate retained under `reports/initial-candidate-receipts/`.
* Native diagram reading order follows named left-to-right groups/rows with full notes equivalents. Native PowerPoint reading-order pane, decorative branding semantics and screen-reader behaviour still require actual viewer testing; no conformance claim.
* Shared renderer emits a nonblocking PyMuPDF deprecation warning; no output was skipped.

## Planned timing and public evidence

* Main1-18:795 seconds, including12 public repository evidence110 seconds. Q&A75 + contingency30 =900 seconds total.
* Exact per-slide seconds:35,25,35,40,25,30,60,40,30,55,40,110,60,50,30,50,45,35.
* Walkthrough retains0-15,15-40,40-70,70-95,95-110 sequence. Offline fallback to native12 within15 seconds is INSIDE110, not added time.
* No live auth, new model call, fake review click or private receipt. Optional appendices use Q&A, not an extra main pitch.
* This is arithmetic and a script plan, not measured delivery or rehearsal.

## Reproduction without uv

* Existing PYTHONPATH supplies python-pptx, PyYAML, Pillow and PyMuPDF; local LibreOffice/pdftoppm and Liberation Sans were retained. CairoSVG and its dependencies were installed with pip into an isolated session target only after the converter import failed. No uv or global installation.
* Complete24 YAML/style/assets, generator, retained structural source, finaliser, reused helper and validator are included inside this candidate.
* The byte-identical 22-slide intake is pinned to the retained `../zava-forward-capstone-rebuild/slide-deck/zava-forward-executive-draft.pptx`, SHA256 `7b25fd7e77d1ed7d9aa908cf3da429bcfeae104beadea5ac90c1e97dbc71fbb4`. Generator, validator and baseline PDF export use this immutable working copy, not the replaceable public delivery copy. With previous22 content still at recorded hashes:

```bash
"""
    content += f'bash "{ROOT}/rebuild.sh"\n'
    content += """```

* `rebuild.sh` invokes actual shared build_deck.py with --source and all24 targets, exports full PDF/JPGs, runs shared property/geometry checks, then validates actual output. Reuses existing PPTX tools plus isolated icon converter target; no uv/model pipeline.
* To rebuild just the retained structural source and complete YAML without regenerating content, use:

```bash
"""
    content += f'cd "{ROOT}"\n'
    content += """export PYTHONDONTWRITEBYTECODE=1
export PYTHONPATH=/home/azureuser/.copilot/session-state/1e0d58ea-7ea9-4c65-93b7-66c89c481468/files/capstone-pptx-tools
export PYTHONPATH="$PYTHONPATH:/home/azureuser/.copilot/session-state/1e0d58ea-7ea9-4c65-93b7-66c89c481468/files/capstone-icon-tools"
tools=/home/azureuser/.copilot/installed-plugins/hve-core/hve-core-all/skills/experimental/powerpoint/scripts
python3 "$tools/build_deck.py" --content-dir content --style content/global/style.yaml --source source/structural-24-slides.pptx --slides 1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16,17,18,19,20,21,22,23,24 --output slide-deck/zava-forward-executive-draft.pptx
python3 generate_content.py --finalise
```

* Then run the export/render/shared checks in rebuild.sh, including the pinned same-exporter22-slide baseline PDF in reports/baseline-comparison.pdf for literal map13 comparison. Integrity validation and handover generation use the retained22-slide intake even after the public delivery copy becomes24 slides. Do not silently validate against a different deck.
* After actual asset inspection, `python3 write_handover.py` records final peer/retention/font/warning reports and this handover without altering the PPTX. Binary zip hashes can change on a rebuild; repeat render and validation for the new binary.

## Complete20-criterion expert coverage and gaps

This is recorded expert artifact assessment, not a new score, panel decision, audience study or floor clearance. The recorded53/100 intake score at f94f53a is historical, not the latest branch score or a fresh assessment. The approved brief's complete-guide contribution/gap check applies to the implemented routes below; local artifact checks add evidence for B6 only, not live performance.

"""
    brief = (ROOT / "research/ux-design-review.md").read_text()
    criterion_table = brief.split("| Criterion | Brief contribution / slide evidence route | Gap not solved by deck refresh |", 1)[1].split("\n\n**Floors", 1)[0]
    criterion_table = criterion_table.replace(
        "Builder render/native file checks and independent reading needed",
        "Local render/package checks completed; native-viewer and independent human reading still needed")
    content += "| Criterion | Implemented route / contribution | Remaining gap |\n" + criterion_table.lstrip("\n") + "\n"
    content += """
* Floors: A>=21/35, B>=21/35, C>=9/15, D>=9/15; no individual criterion1; thought-leadership artifact submitted; required masterclasses attended. A total never overrides a floor.
* SectionD artifact due at least two weeks before panel, reviewed before panel by assigned reviewers; panel receives it five business days beforehand. Deck alone insufficient; companion notes do not prove formal submission, acceptance or reach.
* No fresh score or floor-clearance decision is made. A7 current authority absent/OFF and other genuine evidence gaps must not be disguised to improve optics.

## Human/native-viewer gates and handoff

* Parent independently inspects candidate and source mappings before any publication. Parent alone owns canonical copy, map/architecture previews, index/log/README updates and any PR68 commit/push.
* Native PowerPoint open/font substitution/connector arrowheads, actual reading order, alt/decorative semantics, Tab/Enter focus and slide-link behaviour are untested in the native viewer. Local package validity and LibreOffice images do not establish keyboard/screen-reader accessibility.
* Verify projection/independent reading,200% zoom and descriptive spoken diagram explanation with representative viewers.
* Rehearse actual795/900-second delivery and110-second evidence recovery. Ask listeners to explain oneAPI, narrow proof, human authority and exact next decision; record answers, not assumed comprehension.
* Staff-session/role/source verification, accessibility, semantic/value, restore/SLO/load/unit costs, professional/legal/Finance/resident approvals and any rollout remain separate product/human gates.
* No production, new-spend, legal-currentness, programme delivery or panel acceptance is inferred. Stop at checked working candidate.
"""
    tracked = subprocess.check_output(
        ["git", "-C", str(REPO), "--no-pager", "diff", "--name-only"], text=True).splitlines()
    staged = subprocess.check_output(
        ["git", "-C", str(REPO), "--no-pager", "diff", "--cached", "--name-only"], text=True).splitlines()
    head = subprocess.check_output(
        ["git", "-C", str(REPO), "rev-parse", "HEAD"], text=True).strip()
    parent_owned = {"docs/capstone/README.md", "docs/reference/README.md",
                    "docs/reference/log.md", "docs/architecture/README.md"}
    working_prefix = str(ROOT.relative_to(REPO)) + "/"
    assert all(path in parent_owned or path.startswith(working_prefix)
               for path in tracked) and not staged
    boundary = dict(head=head, tracked_changes=tracked, staged_changes=staged,
                    documentation_intake_commit="f94f53a56c36d38ec36f7f662595912084571380",
                    canonical_intake_hash_unchanged=pkg["canonical_untouched"],
                    allowed_write_root=str(ROOT),
                    approved_ux_brief_read_only=True,
                    parent_owned_expected_doc_changes=[path for path in tracked if path in parent_owned],
                    allowed_working_tree_changes=[path for path in tracked if path.startswith(working_prefix)],
                    no_publication_or_git_write=True,
                    source_24_sha256=hashlib.sha256(
                        (ROOT / "source/structural-24-slides.pptx").read_bytes()).hexdigest(),
                    note="Read-only git checks. Unrelated old untracked outputs not written or removed.")
    (REPORTS / "final-boundary-check.json").write_text(json.dumps(boundary, indent=2) + "\n")
    content += "\n* Read-only generation boundary check: execution HEAD is recorded separately from the fixed documentation intake commit; no staged changes. Tracked changes are limited to this working tree and four parent-owned README/index/log paths. This helper does not publish or write to Git. Pinned22-slide intake hash unchanged. See `reports/final-boundary-check.json`.\n"
    markdown(ROOT / "build-handover.md", "Zava Forward architecture refresh build handover", content)
    markdown(ROOT / "changes/build-execution.md", "Architecture refresh execution log", f"""## Task and boundary

* Complete approved UX brief implemented in this candidate tree only. Brief and inputs remain read-only.
* Source22: `{pkg['canonical_sha256']}`. Final24: `{pkg['sha256']}`.
* No agents/workflows, cloud/auth/model calls, private records, uv, publication or git writes. Parent-authorised isolated CairoSVG installation only.

## Execution and recovery

* Read complete required grounding, approved brief and shared PowerPoint instructions; reused current22 helpers/patterns.
* Generated complete24 YAML/style/content assets and structural source with only two new blank slides. Shared builder --source rebuild includes13 for authorised spacing/frame/counter correction; retained safe22 graphicFrame clear hook.
* Initial executive-source alias mismatch fixed by explicit source aliases before proceeding. This was a build preparation failure, not a fallback to a different deck.
* Fresh actual renders exposed overflowing card copy and tight fixed title/body regions. Shortened labels/title19, increased fixed slack and common peer heights, rerouted gutters; reran full build/export/render/checks.
* Supplemental peer assertion initially required identical body fonts, beyond the approved equal header/font and body-start rule. Corrected the assertion to enforce exact title font/inset/geometry and body geometry, explicitly reporting role-specific body sizes instead of hiding them.
* Supplemental canonical-frame retention found one original18 phrase not identical in baseline YAML (whole-life affordability). Updated generator to retain actual canonical text as well as YAML; reran full build/export/render/checks. All245 original nonmetadata frames now retained.
* Raised shortened source-only refresh footers from9 to10pt and normalised generated report status symbols/currency to ASCII. No essential claim moved to a footer.
* Actual PDF checks detected the increased10pt footer metrics below the previous frame/canvas. Moved these metadata frames up0.02in and increased fixed height to0.21in; reran full pipeline rather than accepting changed-frame overflow.
* Latest `bash rebuild.sh` completed successfully:24-page PDF,24 fresh150dpi JPGs, shared property/geometry and custom actual-output checks.
* Final correction inspection: all four contact sheets plus full-size changed10/13/15/19/20/23/24 and neighbours9/11/12/14/16/18/21/22. Shortened overflowing Blob label after actual image showed a card border through the last line, despite weaker old position attribution passing.
* Supplementary actual-output peer, original-frame retention and explicit-font checks passed. All{len(dispositions)} generic findings retained with class-based per-row dispositions.
* A reporting probe using `python` failed because only `python3` is present; reran with python3. No installation or environment change.

## Final findings

*24/24 relationship-bound note matches; original canvas/master/layout/theme/media preserved; native13 logical content/fonts/structure and original notes retained, exact authorised delta reported.
* Zero native text intersections; all TOP/fixed; intended peers zero drift; architecture remains native, mapped, with20 individual product icons/11 source assets and no composites.
* Shared property0 issues. Shared geometry{len(dispositions)} findings: {severities['warning']} warnings and{severities['info']} info, zero errors. Strict complete expected-sequence glyph bounds TRUE on all387 frames, no exemptions.
* Actual24-page per-frame tokens, all24 counters, links, text/connector contrast, icon provenance/no-crop/aspect, connector routing and claim/status oracle pass.
* See `build-handover.md`, `reports/visual-inspection.md` and `reports/warning-dispositions.json` for exact evidence and limits.
* Candidate complete; native-viewer/accessibility/projection/rehearsal/audience/panel/product gates remain human-owned.
""")
    print("Handover written; final candidate unchanged:", pkg["sha256"])


if __name__ == "__main__":
    main()
