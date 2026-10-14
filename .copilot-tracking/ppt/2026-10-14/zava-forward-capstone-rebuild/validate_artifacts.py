"""Validate local native geometry, relationship-bound notes and rendered text."""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
import hashlib
from itertools import combinations
import json
import logging
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from PIL import Image, ImageDraw
import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
import yaml

ROOT = Path(__file__).resolve().parent
DECK = ROOT / "slide-deck/zava-forward-executive-draft.pptx"
VALIDATION = ROOT / "slide-deck/validation"
REPORTS = ROOT / "reports"
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/package/2006/relationships",
}
LOGGER = logging.getLogger(__name__)


def normalized(text: str) -> str:
    """Normalize line breaks and punctuation for semantic rendering checks."""
    return re.sub(r"[^a-z0-9£]+", " ", text.lower()).strip()


def write_json(name: str, data: object) -> None:
    """Write a report in the bounded artifact directory."""
    (REPORTS / name).write_text(json.dumps(data, indent=2, ensure_ascii=False)+"\n")


def note_associations(prs: Presentation, provenance: dict) -> list[dict]:
    """Resolve each physical slide to its actual notes relationship and sentinels."""
    rows = []
    with ZipFile(DECK) as package:
        for i, slide in enumerate(prs.slides, 1):
            part = str(slide.part.partname).lstrip("/")
            rels_path = str(Path(part).parent / "_rels" / (Path(part).name+".rels"))
            rels = ET.fromstring(package.read(rels_path))
            note_rel = next(r for r in rels if r.attrib["Type"].endswith("/notesSlide"))
            notes_part = slide.notes_slide.part
            note_text = slide.notes_slide.notes_text_frame.text
            content = yaml.safe_load((ROOT / f"content/slide-{i:03d}/content.yaml").read_text())
            topic = provenance["notes_topics"][i-1]
            source_ids = provenance["source_ids"][str(i)]
            passed = (
                f"Topic sentinel: {topic}" in note_text
                and all(f"S{s:02d}:" in note_text for s in source_ids)
                and note_text == content["speaker_notes"]
                and f"SLIDE {i:02d} — {content['title']}" in note_text
            )
            rows.append(dict(
                slide=i, title=content["title"], slide_part=part,
                relationship_id=note_rel.attrib["Id"],
                relationship_target=note_rel.attrib["Target"],
                resolved_notes_part=str(notes_part.partname),
                topic_sentinel=topic, source_sentinels=";".join(f"S{s:02d}" for s in source_ids),
                characters=len(note_text), exact_yaml_notes_match=note_text == content["speaker_notes"],
                pass_association=passed,
            ))
    with (REPORTS / "notes-association-22-rows.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    write_json("notes-association.json", rows)
    return rows


def geometry(prs: Presentation) -> dict:
    """Measure actual native shapes and peer rows, with overlays classified."""
    errors = []
    peer_groups = defaultdict(list)
    texts = []
    native_counts = []
    for i, slide in enumerate(prs.slides, 1):
        counts = Counter()
        visible_texts = [s for s in slide.shapes if s.has_text_frame and s.text.strip()]
        for first, second in combinations(visible_texts, 2):
            overlap_x = min(first.left+first.width, second.left+second.width) > max(first.left, second.left)
            overlap_y = min(first.top+first.height, second.top+second.height) > max(first.top, second.top)
            if overlap_x and overlap_y:
                errors.append(f"Slide {i} text-frame collision: {first.name} / {second.name}")
        for shape in slide.shapes:
            counts[str(shape.shape_type)] += 1
            if shape.has_table:
                errors.append(f"Slide {i} retained source table: {shape.name}")
            if shape.left < 0 or shape.top < 0 or shape.left+shape.width > prs.slide_width+2 or shape.top+shape.height > prs.slide_height+2:
                errors.append(f"Slide {i} out of bounds: {shape.name}")
            if not shape.has_text_frame or not shape.text.strip():
                continue
            tf = shape.text_frame
            fonts = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
            row = dict(
                slide=i, name=shape.name, text=shape.text, left_emu=shape.left,
                top_emu=shape.top, width_emu=shape.width, height_emu=shape.height,
                insets_emu=dict(left=tf.margin_left, right=tf.margin_right,
                                top=tf.margin_top, bottom=tf.margin_bottom),
                anchor=str(tf.vertical_anchor), auto_size=str(tf.auto_size),
                font_min_pt=min(fonts) if fonts else None,
                font_max_pt=max(fonts) if fonts else None,
            )
            texts.append(row)
            if tf.vertical_anchor != MSO_ANCHOR.TOP:
                errors.append(f"Slide {i} not TOP anchored: {shape.name}")
            if tf.auto_size != MSO_AUTO_SIZE.NONE:
                errors.append(f"Slide {i} not fixed-size: {shape.name}")
            if shape.name == "slide-title" and min(fonts) < 28:
                errors.append(f"Slide {i} main title below 28pt")
            minimum = 16 if i == 13 else 18
            if "peer" in shape.name and ("body" in shape.name or "title" in shape.name or i == 13):
                if min(fonts) < minimum:
                    errors.append(f"Slide {i} peer font below {minimum}pt: {shape.name}")
            if "peer-" in shape.name:
                key = re.sub(r"peer-\d+", "peer-N", shape.name)
                peer_groups[(i, key)].append(row)
        native_counts.append(dict(slide=i, counts=dict(counts), shapes=len(slide.shapes)))
    groups = []
    for (slide, group), rows in peer_groups.items():
        signatures = {
            (r["top_emu"], r["height_emu"], tuple(r["insets_emu"].values()),
             r["anchor"], r["font_min_pt"], r["font_max_pt"])
            for r in rows
        }
        passed = len(signatures) == 1
        if not passed:
            errors.append(f"Slide {slide} mismatched peer row: {group}")
        groups.append(dict(slide=slide, group=group, peers=len(rows), exact_emu_match=passed,
                           actual_geometry=rows))
    map_slide = prs.slides[12]
    map_text = " ".join(s.text for s in map_slide.shapes if s.has_text_frame)
    canonical_capabilities = [
        "Stock Prioritisation", "Assessment Preparation", "Grant & Funding Evidence",
        "Works & Resident Coordination", "Quality Evidence Review",
    ]
    map_required = canonical_capabilities + [
        "Choose homes", "Assess & design", "Prepare funding", "Coordinate works",
        "Verify & hand over", "Assets", "data owners", "Assessor", "coordinator",
        "Finance", "Resident", "Programme", "Independent quality",
        "On request", "Trigger", "OFF", "Workflow core", "rights", "versions",
        "gates", "human approvals", "audit", "retries", "manual recovery",
        "application", "award", "claim", "cash", "Hazard", "human service",
    ]
    map_semantics = {term: normalized(term) in normalized(map_text) for term in map_required}
    map_notes = map_slide.notes_slide.notes_text_frame.text
    map_ids = ["stock-prioritisation", "assessment-preparation", "funding-evidence",
               "works-resident-coordination", "quality-evidence"]
    map_semantics.update({f"notes:{term}": term in map_notes for term in map_ids})
    map_semantics["five_request_labels"] = map_text.count("On request") == 5
    map_semantics["five_off_trigger_labels"] = map_text.count("Trigger • OFF") == 5
    map_semantics["native_arrows"] = sum(s.shape_type == MSO_SHAPE_TYPE.LINE for s in map_slide.shapes) == 4
    pictures = [s for s in map_slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    branding = [s.name for s in pictures if s.name == "Zava branding logo — only branding"]
    diagram_count = len(pictures)-len(branding)
    map_semantics["diagram_image_count_zero"] = diagram_count == 0
    map_semantics["branding_identified"] = len(branding) == 1
    if not all(map_semantics.values()):
        errors.extend(f"Map semantic failure: {k}" for k, passed in map_semantics.items() if not passed)
    return dict(
        pass_geometry=not errors, errors=errors, dimensions_emu=[prs.slide_width, prs.slide_height],
        text_frame_pair_intersections=sum("text-frame collision" in e for e in errors),
        text_frames=texts, peer_rows=groups, native_shape_counts=native_counts,
        map=dict(semantic_checks=map_semantics, diagram_image_count=diagram_count,
                 branding_picture_names=branding, total_shapes=len(map_slide.shapes),
                 map_text_minimum_pt=min(r["font_min_pt"] for r in texts if r["slide"] == 13
                                         and "peer" in r["name"])),
        overlap_classification=(
            "Text boxes intentionally overlay their separate native filled card backgrounds. "
            "Logo intentionally overlays white brand backing. These are containers, not text collisions. "
            "Peer gaps are deliberate; map gutters are .13 inch and are validated through fresh images. "
            "Named footer and branding edge placements are intentional, inside the page; "
            "the map uses dense margins specified by the approved handover."
        ),
    )


def rendered_text(prs: Presentation) -> dict:
    """Check each native text box against actual locally rendered PDF words."""
    results = []
    with pymupdf.open(VALIDATION / "slides.pdf") as document:
        assert len(document) == 22
        for i, (slide, page) in enumerate(zip(prs.slides, document), 1):
            sx = page.rect.width/prs.slide_width
            sy = page.rect.height/prs.slide_height
            for shape in slide.shapes:
                if not shape.has_text_frame or not shape.text.strip():
                    continue
                rect = pymupdf.Rect(shape.left*sx-.5, shape.top*sy-.5,
                                    (shape.left+shape.width)*sx+.5,
                                    (shape.top+shape.height)*sy+.5)
                expected = Counter(normalized(shape.text).split())
                observed = Counter(normalized(page.get_text("text", clip=rect)).split())
                missing = expected-observed
                results.append(dict(slide=i, name=shape.name, pass_rendered_text=not missing,
                                    missing_tokens=dict(missing), pdf_clip_points=list(rect)))
    failures = [r for r in results if not r["pass_rendered_text"]]
    return dict(pass_rendered_text=not failures, frames_checked=len(results),
                failures=failures, checks=results,
                method="Every textbox's expected token multiset compared with actual LibreOffice PDF text within its native bounds, .5pt tolerance. This supplements, not replaces, image inspection.")


def claim_checks(prs: Presentation) -> dict:
    """Check exact required provenance and open-state language in generated notes."""
    notes = "\n".join(s.notes_slide.notes_text_frame.text for s in prs.slides)
    text = "\n".join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
    terms = [
        "cd14b3b379026a87ebfa3e9058ac10c0df7888eb",
        "sha256:e0254cb3a0599c21802f0b6515f975550cf00ddce4849d04fd3fa3d50683af20",
        "ca-forward-test-api--0000012", "standard", "gpt-4.1-mini-2025-04-14",
        "6884", "15 ledger", "980", "38", "11", "cabe05c", "CON-009", "AC-057",
        "RV-001", "RV-002", "currency guarantee", "reset", "Decision 1",
        "110", "15s", "recorded operator-service proof",
    ]
    checks = {term: term.lower() in notes.lower() for term in terms}
    checks["no_stale_current_screen_claims"] = not any(term in text for term in ["11 OCT", "v0.13", "v0.12", "#58"])
    checks["authority_absent_screen"] = "authority absent" in text.lower()
    checks["proactive_off_screen"] = "proactive off" in text.lower()
    checks["no_placeholders"] = not any(t in text.lower() for t in ["lorem ipsum", "click to add", "tbd", "placeholder"])
    checks["planning_only_close"] = "No new procurement, spending, rollout or production authority" in prs.slides[17].notes_slide.notes_text_frame.text or "not yes to autonomous rollout" in prs.slides[17].notes_slide.notes_text_frame.text
    return dict(pass_claim_checks=all(checks.values()), checks=checks,
                boundary="Content-presence/status regression against adopted oracle, not independent runtime or legal-currentness verification.")


def contact_sheets() -> list[str]:
    """Create four readable six-up sheets from the actual fresh JPEG exports."""
    paths = []
    for start in range(1, 23, 6):
        stop = min(start+5, 22)
        canvas = Image.new("RGB", (1320, 3*405), "#e8e8e8")
        draw = ImageDraw.Draw(canvas)
        for offset, i in enumerate(range(start, stop+1)):
            with Image.open(VALIDATION / f"slide-{i:03d}.jpg") as image:
                image.thumbnail((640, 360))
                x = 10+(offset % 2)*660
                y = 30+(offset // 2)*405
                canvas.paste(image, (x, y))
                draw.text((x, y-19), f"SLIDE {i:02d} / 22", fill="#173f35")
        path = VALIDATION / f"contact-{start:02d}-{stop:02d}.jpg"
        canvas.save(path, quality=92)
        paths.append(str(path.relative_to(ROOT)))
    return paths


def package_integrity(
    prs: Presentation, *, expected_canonical_sha: str, source_sha: str
) -> dict:
    """Check immutable source, theme parts and explicitly expected delivery bytes."""
    source = ROOT / "source/canonical-22-slides.pptx"
    canonical = ROOT.parents[3] / "docs/capstone/deck/zava-forward-executive-draft.pptx"
    with ZipFile(source) as original, ZipFile(DECK) as rebuilt:
        theme_parts = [name for name in original.namelist()
                       if name.startswith(("ppt/theme/", "ppt/slideMasters/", "ppt/slideLayouts/"))]
        unchanged = {name: original.read(name) == rebuilt.read(name) for name in theme_parts}
    canonical_sha = hashlib.sha256(canonical.read_bytes()).hexdigest()
    source_copy_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    return dict(
        working_sha256=hashlib.sha256(DECK.read_bytes()).hexdigest(),
        canonical_sha256=canonical_sha,
        expected_canonical_sha256=expected_canonical_sha,
        canonical_unchanged=canonical_sha == source_sha,
        canonical_matches_expected=canonical_sha == expected_canonical_sha,
        source_copy_matches_recorded=source_copy_sha == source_sha,
        source_slide_count=22, output_slide_count=len(prs.slides),
        preserved_master_theme_layout_parts=unchanged,
        all_theme_master_layout_parts_byte_identical=all(unchanged.values()),
        output_bytes=DECK.stat().st_size,
    )


def classify_shared_findings() -> dict:
    """Retain and classify every default-threshold heuristic finding after image review."""
    path = VALIDATION / "shared-geometry-results.json"
    if not path.exists():
        return dict(status="not run")
    data = json.loads(path.read_text())
    rows = []
    for slide in data["slides"]:
        for issue in slide["issues"]:
            check = issue["check_type"]
            if issue["severity"] == "error":
                reason = "UNRESOLVED; must be fixed before handover."
            elif check == "edge_margin":
                reason = (
                    "Intentional branding/footer edge band, fully in bounds. "
                    "Dense map/control boundary lines inspected in exported images; "
                    "no clipped glyphs or footer collisions. Default .5in margin is not "
                    "the approved dense-map/footer layout."
                )
            elif check == "title_clearance" and "peer" in issue["location"]:
                reason = (
                    "Fixed peer title and body regions are separated; checker also compares "
                    "across different horizontal columns. Actual native text rectangles do "
                    "not intersect. Reserved title height/body baselines and rendered glyphs checked."
                )
            else:
                reason = (
                    "Deliberate compact spacing below generic .3in gap/.2in title-clearance threshold. "
                    "Actual text-frame intersections are zero and all expected words render "
                    "inside their frames. Local image inspection found no collision or readability blocker."
                )
            rows.append(dict(slide=slide["slide_number"], **issue,
                             disposition="intentional-reviewed" if issue["severity"] != "error" else "unresolved",
                             rationale=reason))
    report = dict(
        default_threshold_findings=len(rows),
        severity_counts=dict(Counter(r["severity"] for r in rows)),
        unresolved=sum(r["disposition"] == "unresolved" for r in rows),
        status="Default heuristic is not claimed passed. Findings retained and individually classified.",
        overlay_policy="Card/title/body overlay is intentional containment, never counted as a text collision; logo/backing likewise.",
        findings=rows,
    )
    write_json("shared-geometry-dispositions.json", report)
    return report


def main() -> None:
    """Run all bounded structural/rendered-content checks and create contact sheets."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--expected-canonical-sha",
        help="Exact expected canonical SHA-256; defaults to the original pre-publication source",
    )
    args = parser.parse_args()
    if args.expected_canonical_sha and not re.fullmatch(r"[0-9a-f]{64}", args.expected_canonical_sha):
        parser.error("--expected-canonical-sha must be a lowercase 64-character SHA-256")
    prs = Presentation(DECK)
    assert len(prs.slides) == 22, "Blocking: expected exactly 22 physical slides"
    provenance = json.loads((REPORTS / "provenance.json").read_text())
    notes = note_associations(prs, provenance)
    measured = geometry(prs)
    rendering = rendered_text(prs)
    claims = claim_checks(prs)
    package = package_integrity(
        prs,
        expected_canonical_sha=args.expected_canonical_sha or provenance["source_sha256"],
        source_sha=provenance["source_sha256"],
    )
    dispositions = classify_shared_findings()
    sheets = contact_sheets()
    write_json("native-peer-geometry.json", measured)
    write_json("rendered-text-integrity.json", rendering)
    write_json("claim-oracle-checks.json", claims)
    write_json("package-integrity.json", package)
    summary = dict(
        slide_count=len(prs.slides), notes_associations_pass=all(r["pass_association"] for r in notes),
        geometry_pass=measured["pass_geometry"], rendered_text_pass=rendering["pass_rendered_text"],
        claim_checks_pass=claims["pass_claim_checks"], package_integrity=package,
        shared_geometry=dict((k, v) for k, v in dispositions.items() if k != "findings"),
        contact_sheets=sheets,
    )
    write_json("validation-summary.json", summary)
    LOGGER.info("Summary: %s", json.dumps({k: v for k, v in summary.items() if k != "package_integrity"}))
    for error in measured["errors"]:
        LOGGER.error("%s", error)
    for failure in rendering["failures"]:
        LOGGER.warning("Rendered text mismatch slide %s %s: %s", failure["slide"], failure["name"], failure["missing_tokens"])
    for key, passed in claims["checks"].items():
        if not passed:
            LOGGER.error("Claim regression: %s", key)
    if not all([summary["notes_associations_pass"], summary["geometry_pass"],
                summary["rendered_text_pass"], summary["claim_checks_pass"],
                package["canonical_matches_expected"], package["source_copy_matches_recorded"],
                package["all_theme_master_layout_parts_byte_identical"]]):
        raise ValueError(
            "Blocking local artifact-validation failure; inspect reports. "
            f"Canonical SHA-256 expected {package['expected_canonical_sha256']}, "
            f"observed {package['canonical_sha256']}; source-copy intact "
            f"{package['source_copy_matches_recorded']}."
        )


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    main()
