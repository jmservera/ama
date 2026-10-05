from __future__ import annotations

import hashlib
import io
import re
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from .domain import SOURCE_STATUSES

MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
MAX_WORKBOOK_ROWS = 1000
REQUIRED_COLUMNS = {
    "candidate_id",
    "property_ref",
    "current_status",
    "owner",
    "is_synthetic",
}
OPTIONAL_COLUMNS = {"locality", "archetype_id", "lead_supplier", "rag_status"}
ALLOWED_COLUMNS = REQUIRED_COLUMNS | OPTIONAL_COLUMNS
CANDIDATE_ID = re.compile(r"^CAND-[A-Z0-9-]{1,40}$")
PROPERTY_REF = re.compile(r"^(SYN-PROP-[A-Z0-9-]+|CAND-[A-Z0-9-]+|HA-[0-9]+|HB/[0-9]+)$")
ALLOWED_VALUES = {
    "owner": frozenset(
        {
            "",
            "BSL-01",
            "GEO-01",
            "Grants",
            "PM-01",
            "PS-01",
            "Programme",
            "QA-01",
            "RC-01",
            "RC-02",
            "RC-03",
            "RLO-01",
            "RLO-02",
            "RLO-03",
            "Resident Liaison",
            "Retrofit Coord",
        }
    ),
    "locality": frozenset({"", "Locality A", "Locality B", "Locality C"}),
    "archetype_id": frozenset({"", *(f"ARC-{number:02}" for number in range(1, 7))}),
    "lead_supplier": frozenset({"", "SUP-F1", "SUP-F2", "SUP-HP", "SUP-HT"}),
    "rag_status": frozenset({"", "Amber", "Green", "Red"}),
}


@dataclass
class WorkbookPreview:
    digest: str
    source_name: str
    accepted: list[dict[str, Any]]
    rejected: list[dict[str, Any]]
    warnings: list[str]


def parse_workbook(content: bytes, source_name: str, existing_ids: set[str]) -> WorkbookPreview:
    if not content:
        raise ValueError("The uploaded workbook is empty.")
    if len(content) > MAX_WORKBOOK_BYTES:
        raise ValueError("Workbook exceeds the 5 MB local prototype limit.")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("Could not read this file as an Excel workbook (.xlsx).") from exc

    try:
        if "cases" not in workbook.sheetnames:
            raise ValueError("Workbook must contain a worksheet named 'cases'.")
        sheet = workbook["cases"]
        if sheet.max_row is None or sheet.max_row < 2:
            raise ValueError("The 'cases' worksheet must include headers and at least one row.")
        if sheet.max_row - 1 > MAX_WORKBOOK_ROWS:
            raise ValueError("Workbook exceeds the 1,000-row local prototype limit.")

        raw_headers = [cell.value for cell in sheet[1]]
        headers = [
            str(value).strip().casefold() if value is not None else ""
            for value in raw_headers
        ]
        if len(headers) != len(set(headers)):
            raise ValueError("Workbook contains duplicate column headers.")
        missing = sorted(REQUIRED_COLUMNS - set(headers))
        unknown = sorted(set(headers) - ALLOWED_COLUMNS)
        if "" in headers:
            unknown.append("<blank header>")
        if missing:
            raise ValueError(f"Workbook is missing required columns: {', '.join(missing)}.")
        unknown_header_issue = (
            f"Workbook contains unmapped columns: {', '.join(unknown)}."
            if unknown
            else None
        )

        accepted: list[dict[str, Any]] = []
        rejected: list[dict[str, Any]] = []
        seen: set[str] = set()
        warnings: list[str] = []
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=2, values_only=False), start=2
        ):
            values: dict[str, Any] = {}
            formula_found = False
            for header, cell in zip(headers, cells):
                if not header:
                    continue
                if cell.data_type == "f":
                    formula_found = True
                values[header] = cell.value
            if not any(value is not None and str(value).strip() for value in values.values()):
                continue

            issues: list[str] = []
            if unknown_header_issue:
                issues.append(unknown_header_issue)
            candidate_id = str(values.get("candidate_id") or "").strip().upper()
            property_ref = str(values.get("property_ref") or "").strip().upper()
            is_synthetic = str(values.get("is_synthetic") or "").strip().casefold()
            if formula_found:
                issues.append("Formula cells are not accepted.")
            if not CANDIDATE_ID.fullmatch(candidate_id):
                issues.append("candidate_id has an unsupported format.")
            if not PROPERTY_REF.fullmatch(property_ref):
                issues.append("property_ref must be a synthetic reference.")
            if is_synthetic != "true":
                issues.append("is_synthetic must be true for every row.")
            if not str(values.get("current_status") or "").strip():
                issues.append("current_status is required.")
            elif str(values["current_status"]).strip() not in SOURCE_STATUSES:
                issues.append("current_status has an unmapped value.")
            for field, allowed in ALLOWED_VALUES.items():
                value = str(values.get(field) or "").strip()
                if value not in allowed:
                    issues.append(f"{field} has an unmapped value.")
            if candidate_id in seen:
                issues.append("candidate_id is duplicated in this workbook.")
            if candidate_id in existing_ids:
                issues.append("candidate_id already exists; imports never overwrite cases.")
            if candidate_id:
                seen.add(candidate_id)

            record = {
                "candidate_id": candidate_id,
                "property_ref": property_ref,
                "current_status": str(values.get("current_status") or "").strip(),
                "owner": str(values.get("owner") or "").strip(),
                "locality": str(values.get("locality") or "").strip(),
                "archetype_id": str(values.get("archetype_id") or "").strip(),
                "lead_supplier": str(values.get("lead_supplier") or "").strip(),
                "rag_status": str(values.get("rag_status") or "").strip(),
                "is_synthetic": True,
                "source_row": row_number,
            }
            if issues:
                rejected.append(
                    {"source_row": row_number, "candidate_id": candidate_id, "issues": issues}
                )
            else:
                accepted.append(record)

        if not accepted and not rejected:
            raise ValueError("No non-empty case rows were found in the workbook.")
        if rejected:
            warnings.append(
                f"{len(rejected)} row(s) need human correction; no rejected row will be imported."
            )
        return WorkbookPreview(
            digest=hashlib.sha256(content).hexdigest(),
            source_name=source_name[:160],
            accepted=accepted,
            rejected=rejected,
            warnings=warnings,
        )
    finally:
        workbook.close()
